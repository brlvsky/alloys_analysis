"""Пакетный режим: графики, summary.xlsx, report.html."""
from pathlib import Path

import pytest
from openpyxl import load_workbook

from psd_lab.core.report import run_batch
from psd_lab.core.settings import Settings, format_windows, parse_windows

RAW = Path(__file__).resolve().parent.parent / "data" / "raw"


def test_batch(tmp_path):
    if not (RAW / "N_C_.xls").exists():
        pytest.skip("нет исходных файлов")
    res = run_batch(RAW, tmp_path, Settings(), log=lambda *_: None)
    assert len(res.samples) == 12
    assert (tmp_path / "compare.png").exists()
    assert len(list(tmp_path.rglob("*.png"))) == 13
    wb = load_workbook(tmp_path / "summary.xlsx")
    assert wb.sheetnames == ["Сводка", "Кривые", "Флаги", "Популяции", "Окна печати", "Выход годного",
                             "Поверхность", "Упаковка", "Кинетика"]
    assert wb["Сводка"].max_row == 13
    html = (tmp_path / "report.html").read_text(encoding="utf-8")
    # 13 графиков распределений + 6 популяций (многомодальные) + полосы окон печати + кинетика
    assert html.count("data:image/png;base64,") == 13 + 6 + 1 + 1
    for sec in ("Популяции частиц", "окна СЛС и СЭЛС", "Выход годного", "Удельная поверхность",
                "плотности упаковки", "Кинетика помола", "агломерация"):
        assert sec in html
    assert html.count("Допущения:") >= 3
    assert "П/С +0,5Y2O3" in html and "ОШИБКА" in html
    assert "src='http" not in html and 'src="http' not in html  # самодостаточный файл


def test_windows_text_roundtrip():
    w = parse_windows("<15; 15-45; 15–53; 45-105; >53")
    assert w == [[None, 15], [15, 45], [15, 53], [45, 105], [53, None]]
    assert parse_windows(format_windows(w)) == w
    with pytest.raises(ValueError):
        parse_windows("45-15")


def test_settings_roundtrip(tmp_path):
    s = Settings(lang="ru", bin_um=2.0)
    s.add_recent("C:/данные/Расчет.xlsx")
    s.save(tmp_path / "settings.json")
    s2 = Settings.load(tmp_path / "settings.json")
    assert s2.lang == "ru" and s2.bin_um == 2.0 and s2.recent_files == ["C:/данные/Расчет.xlsx"]


def test_batch_with_database(tmp_path):
    if not (RAW / "N_C_.xls").exists():
        pytest.skip("нет исходных файлов")
    dbf = tmp_path / "база.sqlite"
    run_batch(RAW, tmp_path / "out", Settings(), log=lambda *_: None, db_file=dbf)
    html = (tmp_path / "out" / "report.html").read_text(encoding="utf-8")
    assert "11. Партии (база данных" in html and "уточнить" in html
    from psd_lab.core import db

    conn = db.connect(dbf)
    assert conn.execute("SELECT COUNT(*) FROM measurements").fetchone()[0] == 16
    run_batch(RAW, tmp_path / "out2", Settings(), log=lambda *_: None, db_file=dbf)   # повтор — без дублей
    assert conn.execute("SELECT COUNT(*) FROM measurements").fetchone()[0] == 16
