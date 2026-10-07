"""Калькулятор шихты (версия 1.1): эталонный пример коллеги из TASK_v1.1.md, часть 3.

Допуски ТЗ: ±0,0005 п.п. для мас.%, ±0,001 г для масс.
"""
import pytest

from psd_lab.core import charge
from psd_lab.core.charge import Additive, ChargeError
from psd_lab.core.elements import masses, molar_mass, parse_formula

PCT = 0.0005 + 1e-9   # допуск, п.п. (+ запас на представление чисел)
G = 0.001 + 1e-9      # допуск, г
BASE = "50Ti-44Al-4.9Nb-1Mo-0.1B"
WT31 = {"Ti": 57.9110, "Al": 28.7259, "Nb": 11.0153, "Mo": 2.3217, "B": 0.0262}


def wt_of(text=BASE, overrides=None):
    comp = charge.parse_composition(text)
    return charge.to_wt(comp.values, masses(overrides))


def close(d, ref, tol):
    for k, v in ref.items():
        assert d[k] == pytest.approx(v, abs=tol), (k, d[k], v)


# ---------------------------------------------------------------- 2.1 атомные массы и формулы
def test_reference_masses_and_formulas():
    m = masses()
    for el, v in {"Ti": 47.867, "Al": 26.9815385, "Nb": 92.90637, "Mo": 95.95, "B": 10.81, "C": 12.011,
                  "Si": 28.085, "Ce": 140.116, "O": 15.999, "Y": 88.90584}.items():
        assert m[el] == v
    for el in "H B C N O Al Si Ti V Cr Mn Fe Co Ni Cu Y Zr Nb Mo Ce Hf Ta W".split():
        assert el in m
    assert parse_formula("Y2O3") == {"Y": 2, "O": 3}
    assert parse_formula("TiB2") == {"Ti": 1, "B": 2}
    assert parse_formula("Ca(OH)2") == {"Ca": 1, "O": 2, "H": 2}
    assert molar_mass("CeO2") == pytest.approx(140.116 + 2 * 15.999)
    assert molar_mass("Nb2O5") == pytest.approx(2 * 92.90637 + 5 * 15.999)
    with pytest.raises(ValueError):
        parse_formula("Xx2O3")


# ---------------------------------------------------------------- 3.1 базовый состав
def test_31_base_composition():
    comp = charge.parse_composition(BASE)
    assert comp.values == {"Ti": 50, "Al": 44, "Nb": 4.9, "Mo": 1, "B": 0.1}
    assert charge.sum_xm(comp.values, masses()) == pytest.approx(4132.81, abs=0.005)
    w = charge.to_wt(comp.values, masses())
    close(w, WT31, PCT)
    assert sum(w.values()) == pytest.approx(100, abs=1e-9)
    g = charge.weigh(w, 200)
    close(g, {"Ti": 115.822, "Al": 57.452, "Nb": 22.031, "Mo": 4.643, "B": 0.052}, G)
    # туда и обратно
    back = charge.to_at(w, masses())
    close(back, comp.values, 1e-9)


def test_31_base_remainder_and_decimal_comma():
    a = charge.parse_composition("Ti-44Al-4,9Nb-1Mo-0,1B")
    b = charge.parse_composition("50 Ti – 44 Al – 4.9 Nb – 1 Mo – 0.1 B")
    assert a.values["Ti"] == pytest.approx(50) and a.base == "Ti"
    close(b.values, a.values, 1e-12)


# ---------------------------------------------------------------- 3.2 загрузки
def test_32_loads():
    w = wt_of()
    full = charge.plan_loads(1500, 200, mode="full")
    assert (full.n, full.per_load_g, full.made_g, full.excess_g) == (8, 200, 1600, 100)
    close(charge.purchase(w, full), {"Ti": 926.575, "Al": 459.615, "Nb": 176.245, "Mo": 37.147, "B": 0.419}, G)
    eq = charge.plan_loads(1500, 200, mode="equal")
    assert eq.n == 8 and eq.per_load_g == pytest.approx(187.5) and eq.excess_g == pytest.approx(0)
    close(charge.weigh(w, eq.per_load_g), {"Ti": 108.583, "Al": 53.861, "Nb": 20.654, "Mo": 4.353, "B": 0.049}, G)
    assert "шаров" in eq.note
    close(charge.weigh(w, 1500), {"Ti": 868.664, "Al": 430.889, "Nb": 165.229, "Mo": 34.825, "B": 0.392}, G)
    # запас на потери и масса шаров
    r = charge.plan_loads(1500, 200, reserve_pct=10, mode="full", balls_ratio=10)
    assert r.n == 9 and r.reserve_g == pytest.approx(150) and r.excess_g == pytest.approx(150)
    assert r.balls_g == pytest.approx(2000)


# ---------------------------------------------------------------- 3.3 добавки
C = Additive("C", 0.5, "at")
SI = Additive("Si", 0.5, "at")
CEO2 = Additive("CeO2", 1.2, "mol")


def test_33_each_additive():
    comp = charge.parse_composition(BASE)
    v = charge.variants(comp, [C, SI, CEO2], mode="each")
    assert [x.name for x in v][0] == "Базовый состав" and len(v) == 4
    close(v[1].w, {"Ti": 57.5816, "Al": 28.8511, "Nb": 11.0633, "Mo": 2.3318, "B": 0.0263, "C": 0.1459}, PCT)
    close(v[2].w, {"Ti": 57.4694, "Al": 28.7948, "Nb": 11.0417, "Mo": 2.3272, "B": 0.0262, "Si": 0.3406}, PCT)
    ce = v[3]
    close(ce.w, {"Ti": 54.5530, "Al": 27.7257, "Nb": 10.6317, "Mo": 2.2408, "B": 0.0252, "CeO2": 4.8235}, PCT)
    close(ce.elements_wt, {"Ce": 3.9267, "O": 0.8967}, PCT)
    assert ce.oxygen["CeO2"] == pytest.approx(0.8967, abs=PCT)
    assert any(lv == "WARN" and "кислород" in t for lv, t in ce.flags)          # O > 0,2 мас.%
    assert not v[1].flags
    # рецепты на 200 г
    close(charge.weigh(v[1].w, 200), {"Ti": 115.163, "Al": 57.702, "Nb": 22.127, "Mo": 4.664, "B": 0.053,
                                      "C": 0.292}, G)
    close(charge.weigh(ce.w, 200), {"Ti": 109.106, "Al": 55.451, "Nb": 21.263, "Mo": 4.482, "B": 0.050,
                                    "CeO2": 9.647}, G)


def test_33_all_together_and_mass_percent():
    comp = charge.parse_composition(BASE)
    allv = charge.variants(comp, [C, SI, CEO2], mode="all")[-1]
    close(allv.w, {"Ti": 53.7846, "Al": 27.9070, "Nb": 10.7013, "Mo": 2.2555, "B": 0.0254, "C": 0.1412,
                   "Si": 0.3301, "CeO2": 4.8550}, PCT)
    assert allv.oxygen_total == pytest.approx(0.9026, abs=PCT)
    both = charge.variants(comp, [C, SI, CEO2], mode="both")
    assert len(both) == 5 and both[-1].name == "Все добавки вместе"
    inst = charge.variants(comp, [Additive("CeO2", 1.2, "wt")], mode="each")[1]
    close(inst.w, {"Ti": 56.7110, "Al": 28.7259, "Nb": 11.0153, "Mo": 2.3217, "B": 0.0262, "CeO2": 1.2}, PCT)
    over = charge.variants(comp, [Additive("CeO2", 1.2, "wt", mode="over")], mode="each")[1]
    close(over.w, {"Ti": 57.2160, "Al": 28.3812, "Nb": 10.8831, "Mo": 2.2938, "B": 0.0258, "CeO2": 1.2}, PCT)


def test_33_density_via_m7():
    comp = charge.parse_composition(BASE)
    v = charge.variants(comp, [CEO2], mode="each")
    for x in v:
        assert x.density is not None and 3.5 < x.density.rho < 4.5
    assert v[1].density.rho > v[0].density.rho                 # CeO₂ тяжелее матрицы


# ---------------------------------------------------------------- 3.4 лигатура
def test_34_ligature_computed_and_given():
    w = wt_of()
    lig = charge.ligature_computed(w, {"Al"})
    assert lig.lig_pct == pytest.approx(71.274, abs=0.0005)
    assert lig.pure["Al"] == pytest.approx(28.726, abs=0.0005)
    close(lig.lig_comp, {"Ti": 81.2511, "Nb": 15.4548, "Mo": 3.2574, "B": 0.0367}, PCT)
    assert 200 * lig.lig_pct / 100 == pytest.approx(142.548, abs=G)
    # заданная лигатура того же состава → тот же результат, отклонений нет
    g = charge.ligature_given(w, lig.lig_comp, {"Al"})
    assert g.ok and g.lig_pct == pytest.approx(71.274, abs=0.001)
    assert max(abs(d) for d in g.deviation.values()) < 1e-6
    # лигатура из Excel коллеги (не соответствует составу) → красный флаг и отклонения
    bad = charge.ligature_given(w, {"Ti": 80.85, "Nb": 15.87, "Mo": 3.24, "B": 0.04}, {"Al"})
    assert not bad.ok and any("точно попасть" in t for _, t in bad.flags)
    # если разрешить добавлять чистые Ti, Nb, Mo, B — попадание точное
    fix = charge.ligature_given(w, {"Ti": 80.85, "Nb": 15.87, "Mo": 3.24, "B": 0.04}, {"Al", "Ti", "Nb", "Mo", "B"})
    assert fix.ok and max(abs(d) for d in fix.deviation.values()) < 1e-4 and fix.lig_pct > 60


# ---------------------------------------------------------------- 3.5 ручной расчёт коллеги
def test_35_hand_calculation():
    ov = {"Ti": 47.87, "Al": 26.98, "Nb": 92.91}
    comp = charge.parse_composition(BASE)
    assert charge.sum_xm(comp.values, masses(ov)) == pytest.approx(4132.91, abs=0.005)
    close(wt_of(overrides=ov), {"Ti": 57.9132, "Al": 28.7236, "Nb": 11.0155, "Mo": 2.3216, "B": 0.0262}, PCT)
    # бор, округлённый до 0,03 мас.% → предупреждение
    rounded = charge.parse_composition("Ti-28.73Al-11.02Nb-2.32Mo-0.03B", basis="wt")
    warns = charge.rounding_warnings(rounded, masses())
    assert any(t.startswith("B ") for t in warns)
    assert charge.weigh(rounded.values, 200)["B"] == pytest.approx(0.06)
    exact = charge.parse_composition("Ti-28.7259Al-11.0153Nb-2.3217Mo-0.0262B", basis="wt")
    assert not charge.rounding_warnings(exact, masses())


# ---------------------------------------------------------------- 3.6 некорректный ввод
def test_36_errors():
    with pytest.raises(ChargeError, match="сумма"):
        charge.parse_composition("50Ti-44Al-4.9Nb-1Mo")              # 99,9 без основы
    with pytest.raises(ChargeError, match="Xx"):
        charge.parse_composition("50Ti-44Xx")
    comp = charge.parse_composition(BASE)
    with pytest.raises(ChargeError, match="единицы"):
        charge.variants(comp, [Additive("CeO2", 1.2, None)])
    with pytest.raises(ChargeError, match="не определён"):
        charge.variants(comp, [Additive("CeO2", 1.2, "at")])
    with pytest.raises(ChargeError, match="только"):
        charge.variants(comp, [Additive("C", 60, "at")])              # больше, чем Ti
    with pytest.raises(ChargeError, match="ёмкость"):
        charge.plan_loads(1500, 0)
    with pytest.raises(ChargeError, match="ёмкость"):
        charge.plan_loads(1500, -5)


# ---------------------------------------------------------------- вывод и чистота
def test_formatting_and_purity():
    assert charge.fmt_pct(0.0262) == "0,02620" and charge.fmt_pct(0.8967) == "0,8967"
    assert charge.fmt_pct(57.911) == "57,91"
    assert charge.fmt_g(0.052) == "0,052" and charge.fmt_g(115.822) == "115,82"
    w = wt_of()
    g = charge.weigh(w, 200, {"Ti": 99.5})
    assert g["Ti"] == pytest.approx(115.822 / 0.995, abs=G)
    with pytest.raises(ChargeError):
        charge.weigh(w, 200, {"Ti": 0})


def test_table_input():
    comp = charge.composition_from_table([("Ti", None, "at"), ("Al", 44, "at"), ("Nb", "4,9", "at"),
                                          ("Mo", 1, "at"), ("B", 0.1, "at")])
    close(comp.values, {"Ti": 50, "Al": 44, "Nb": 4.9, "Mo": 1, "B": 0.1}, 1e-12)
    with pytest.raises(ChargeError, match="одни единицы"):
        charge.composition_from_table([("Ti", None, "at"), ("Al", 28.7, "wt")])


# ---------------------------------------------------------------- рецепт целиком, экспорт, база
def test_recipe_roundtrip_and_calculate():
    import json

    r = charge.EXAMPLE
    res = charge.calculate(r)
    assert [v.name for v in res.variants][0] == "Базовый состав" and len(res.variants) == 5
    assert res.plan.n == 8 and res.plan.per_load_g == 200
    close(res.loads(res.variants[0]), {"Ti": 115.822, "Al": 57.452, "Nb": 22.031, "Mo": 4.643, "B": 0.052}, G)
    close(res.purchase(res.variants[0]), {"Ti": 926.575, "Al": 459.615, "Nb": 176.245, "Mo": 37.147,
                                          "B": 0.419}, G)
    assert res.ligature.lig_pct == pytest.approx(72.093, abs=0.001)   # лигатура варианта «все вместе»
    back = charge.Recipe.from_json(json.loads(json.dumps(r.to_json())))
    assert back.composition == r.composition and len(back.additives) == 3
    assert back.additives[2].component == "CeO2" and back.additives[2].unit == "mol"
    r2 = charge.calculate(back)
    close(r2.main.w, res.main.w, 1e-12)
    with pytest.raises(ChargeError, match="незнакомые поля"):
        charge.Recipe.from_json({"composition": "Ti-44Al", "неизвестное": 1})


def test_exports(tmp_path):
    import docx
    from openpyxl import load_workbook

    from psd_lab.core import charge_report

    res = charge.calculate(charge.EXAMPLE)
    p = charge_report.write_xlsx(res, tmp_path / "Шихта.xlsx", variant=res.variants[0])
    wb = load_workbook(p)
    assert wb.sheetnames == ["Состав", "Варианты", "Загрузки", "Закупка", "Лигатура", "Бланк навесок"]
    assert wb["Состав"]["C2"].value == pytest.approx(57.9110, abs=PCT)      # мас.% Ti
    assert wb["Состав"]["D2"].value == pytest.approx(115.822, abs=0.001)    # на 200 г
    assert wb["Закупка"]["B2"].value == pytest.approx(926.575, abs=0.01)    # Ti на 8 загрузок
    blank = [c.value for row in wb["Бланк навесок"].iter_rows() for c in row]
    assert "Загрузка 1 из 8 — 200,00 г" in blank and "Фактически взвешено, г" in blank

    d = charge_report.write_blank_docx(res, tmp_path / "Бланк.docx")
    text = "\n".join(p.text for p in docx.Document(str(d)).paragraphs)
    assert "Бланк навесок шихты" in text and "Загрузка 8 из 8" in text
    cells = [c.text for t in docx.Document(str(d)).tables for row in t.rows for c in row.cells]
    assert "Точность весов" in cells and "0,001 г (аналитические)" in cells
    # одна таблица на всю партию вместо листа на загрузку
    d2 = charge_report.write_blank_docx(res, tmp_path / "Бланк-партия.docx", per_load=False)
    text2 = "\n".join(p.text for p in docx.Document(str(d2)).paragraphs)
    assert "Вся партия" in text2 and "Загрузка 1 из 8" not in text2


def test_database_recipes(tmp_path):
    import sqlite3

    from psd_lab.core import db

    # старая база версии 1.0 (без таблицы рецептов) открывается без потери данных
    f = tmp_path / "старая.sqlite"
    con = sqlite3.connect(f)
    con.executescript("CREATE TABLE batches(id INTEGER PRIMARY KEY, name TEXT UNIQUE NOT NULL, alloy TEXT, "
                      "additive TEXT, additive_wt_pct REAL, state TEXT, route TEXT, composition TEXT, "
                      "density_measured REAL, notes TEXT);")
    con.execute("INSERT INTO batches(name) VALUES('П/С +0,5Y2O3')")
    con.commit()
    con.close()
    conn = db.connect(f)
    assert [b["name"] for b in db.batches(conn)] == ["П/С +0,5Y2O3"]
    assert conn.execute("PRAGMA user_version").fetchone()[0] == db.SCHEMA_VERSION
    bid = db.get_or_create_batch(conn, "П/С +0,5Y2O3")
    rid = db.charge_save(conn, charge.EXAMPLE.to_json(), "Пример", bid)
    rows = db.charge_list(conn, bid)
    assert len(rows) == 1 and rows[0]["name"] == "Пример" and rows[0]["version"] == "1.1"
    got = charge.Recipe.from_json(db.charge_get(conn, rid))
    assert got.composition == charge.EXAMPLE.composition
    db.charge_delete(conn, rid)
    assert db.charge_list(conn, bid) == []
