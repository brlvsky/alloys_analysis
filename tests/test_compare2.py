"""Сравнение двух образцов «до и после» (ядро)."""
from pathlib import Path

import numpy as np
import pytest

from psd_lab.core import compare2
from psd_lab.core.model import Sample

RAW = Path(__file__).resolve().parent.parent / "data" / "raw"


def lognormal(d50, sg, sizes):
    from scipy.stats import norm

    return Sample(name=f"d50={d50}", size_um=sizes, cum_pct=100 * norm.cdf(np.log(sizes / d50) / np.log(sg)))


def test_identical_samples_zero_difference():
    x = np.linspace(1, 200, 200)
    s = lognormal(30, 1.6, x)
    d = compare2.compare(s, s)
    assert np.allclose(d.delta, 0)
    assert all(abs(m.delta) < 1e-12 for m in d.metrics)
    assert "в пределах разброса повторов" in compare2.summary_text(d)


def test_finer_sample_positive_delta_and_different_grids():
    a = lognormal(40, 1.5, np.linspace(1, 300, 300))          # шаг 1 мкм
    b = lognormal(20, 1.5, np.arange(1.5, 301.5, 2.0))        # шаг 2 мкм, другие узлы, мельче
    d = compare2.compare(a, b)
    assert len(d.grid) > 300                                  # общая сетка — объединение
    assert (d.delta >= -1e-9).all()                           # B мельче во всём диапазоне
    x, dv = d.max_delta
    assert dv > 30 and 15 < x < 60
    d50 = next(m for m in d.metrics if m.name == "d50")
    assert d50.a == pytest.approx(40, rel=0.01) and d50.b == pytest.approx(20, rel=0.01)
    assert d50.rel_pct == pytest.approx(-50, abs=1)
    assert "мельче" in compare2.summary_text(d)


def test_real_pair_and_quality_warning():
    if not (RAW / "N_C_.xls").exists():
        pytest.skip("нет исходных файлов")
    from psd_lab.core import load
    from psd_lab.core.metrics import average_repeats

    by = {s.name: s for s in average_repeats(load(RAW / "N_C_.xls"))}
    d = compare2.compare(by["N/C"], by["N/C+si"])
    rows = {r[0]: r for r in compare2.table_rows(d)}
    assert rows["d50, мкм"][1:4] == ["19,64", "13,31", "−6,33"]
    assert any("N/C+si" in w for w in d.warnings)             # у N/C+si отрицательная обскурация


def test_pair_in_reports(tmp_path):
    if not (RAW / "N_C_.xls").exists():
        pytest.skip("нет исходных файлов")
    import docx
    from openpyxl import load_workbook

    from psd_lab.core.report import run_batch
    from psd_lab.core.settings import Settings

    st = Settings()
    st.compare_pair = ["N/C (M4246+M4247)", "П/С +0,5Y2O3 (M4341)"]
    run_batch(RAW, tmp_path, st, log=lambda *_: None)
    html = (tmp_path / "report.html").read_text(encoding="utf-8")
    assert "До и после: A — N/C (M4246+M4247), B — П/С +0,5Y2O3 (M4341)" in html
    assert "наибольшее расхождение" in html
    text = "\n".join(p.text for p in docx.Document(str(tmp_path / "report.docx")).paragraphs)
    assert "До и после" in text
    ws = load_workbook(tmp_path / "summary.xlsx")["До и после"]
    assert ws["A1"].value == "Показатель" and ws["B3"].value == pytest.approx(19.64, abs=0.01)
