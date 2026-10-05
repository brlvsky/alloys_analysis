"""М1: база данных «структура — свойства»."""
from pathlib import Path

import pytest
from openpyxl import load_workbook

from psd_lab.core import db, file_sha1, load

RAW = Path(__file__).resolve().parent.parent / "data" / "raw"


def test_parse_name():
    p = db.parse_name("П/С +0,5Y2O3")
    assert (p["state"], p["additive"], p["additive_wt_pct"]) == ("П/С", "Y2O3", 0.5)
    p = db.parse_name("N/C+si")
    assert (p["state"], p["additive"], p["additive_wt_pct"]) == ("N/C", "Si", None)
    p = db.parse_name("П/С+1.5 Si")
    assert (p["state"], p["additive"], p["additive_wt_pct"]) == ("П/С", "Si", 1.5)
    p = db.parse_name("N/C")
    assert (p["state"], p["additive"]) == ("N/C", None)
    assert "уточнить" in p["notes"]
    p = db.parse_name("6ч")
    assert p["state"] is None and "6 ч" in p["notes"]


def test_energy_density():
    # P = 200 Вт, v = 1000 мм/с, h = 100 мкм, t = 30 мкм → 200 / (1000·0,1·0,03) = 66,7 Дж/мм³
    assert db.energy_density(200, 1000, 100, 30) == pytest.approx(66.667, rel=1e-4)
    assert db.energy_density(200, None, 100, 30) is None


@pytest.fixture(scope="module")
def conn(tmp_path_factory):
    if not (RAW / "N_C_.xls").exists():
        pytest.skip("нет data/raw")
    c = db.connect(tmp_path_factory.mktemp("db") / "Тестовая база.sqlite")
    for f in ("N_C_.xls", "TANMB_.xls", "Расчет.xlsx"):
        db.import_samples(c, load(RAW / f), file_sha1(RAW / f))
    return c


def test_import_and_no_duplicates(conn):
    n = conn.execute("SELECT COUNT(*) FROM measurements").fetchone()[0]
    assert n == 16                                   # 8 + 5 + 3 исходных измерения
    assert len(db.batches(conn)) == 12               # повторы N/C попадают в одну партию
    again = db.import_samples(conn, load(RAW / "N_C_.xls"), file_sha1(RAW / "N_C_.xls"))
    assert again == 0
    assert conn.execute("SELECT COUNT(*) FROM measurements").fetchone()[0] == 16
    b = conn.execute("SELECT * FROM batches WHERE name='П/С +0,5Y2O3'").fetchone()
    assert (b["state"], b["additive"], b["additive_wt_pct"]) == ("П/С", "Y2O3", 0.5)
    pts = conn.execute("SELECT COUNT(*) FROM psd_points").fetchone()[0]
    assert pts > 16 * 90


def test_metrics_and_modules(conn):
    b = conn.execute("SELECT id FROM batches WHERE name='П/С +0,5Y2O3'").fetchone()["id"]
    (m,) = db.batch_measurements(conn, b)
    assert m["d50"] == pytest.approx(21.87, abs=0.1)
    assert m["fine_pop_pct"] == pytest.approx(52, abs=5)
    assert m["obscuration"] == 33
    mods = conn.execute("SELECT modules FROM psd_metrics WHERE measurement_id=?", (m["id"],)).fetchone()[0]
    assert "deconv" in mods and "surface" in mods


def test_crud_and_structure_property(conn):
    bid = conn.execute("SELECT id FROM batches WHERE name='П/С +0,5Y2O3'").fetchone()["id"]
    bid2 = conn.execute("SELECT id FROM batches WHERE name='N/C'").fetchone()["id"]
    bid3 = conn.execute("SELECT id FROM batches WHERE name='П/С +0,5C'").fetchone()["id"]
    with pytest.raises(ValueError):
        db.insert(conn, "chem", {"O_ppm": "много"}, bid)
    for b, o in ((bid, "1200"), (bid2, "900,5")):
        db.insert(conn, "chem", {"O_ppm": o}, b)
    pj = db.insert(conn, "print_jobs", {"process": "СЭЛС", "power_W": "600", "speed_mm_s": "4000",
                                        "hatch_um": "100", "layer_um": "70", "rel_density_pct": "99,1"}, bid)
    assert db.get(conn, "print_jobs", pj)["energy_density_J_mm3"] == pytest.approx(600 / (4000 * 0.1 * 0.07))
    db.insert(conn, "mech_tests", {"print_job_id": pj, "test_type": "растяжение", "uts_MPa": "850"})
    assert len(db.rows(conn, "mech_tests", bid)) == 1
    pts = db.xy(conn, "psd.d50", "chem.O_ppm")
    assert [p[0] for p in pts] == ["N/C", "П/С +0,5Y2O3"]
    assert db.regression([p[1] for p in pts], [p[2] for p in pts]).r2 is None     # мало точек
    db.insert(conn, "chem", {"O_ppm": "1500"}, bid3)
    pts = db.xy(conn, "psd.d50", "chem.O_ppm")
    reg = db.regression([p[1] for p in pts], [p[2] for p in pts])
    assert reg.n == 3 and reg.r2 is not None
    keys = [k for k, _ in db.property_fields(conn)]
    assert "frac.15–53" in keys and "mech.uts_MPa" in keys
    db.delete(conn, "print_jobs", pj)               # каскадом удаляются и испытания
    assert db.rows(conn, "mech_tests", bid) == []


def test_export(conn, tmp_path):
    p = db.export_xlsx(conn, tmp_path / "база.xlsx")
    wb = load_workbook(p)
    assert wb.sheetnames[:3] == ["Партии", "Измерения", "Метрики"]
    assert wb["Измерения"].max_row == 17
