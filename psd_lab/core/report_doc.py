"""Отчёт как список блоков (заголовок, абзац, таблица, рисунок, примечание…) и два способа его вывести:
самодостаточный HTML и документ Word (.docx). Содержание собирается один раз (report.build_blocks),
поэтому HTML и Word всегда совпадают по составу.
"""
from __future__ import annotations

import html
import io
from dataclasses import dataclass, field
from pathlib import Path


# ---------------------------------------------------------------- блоки
@dataclass
class Heading:
    text: str
    level: int = 2          # 1 — заголовок отчёта, 2 — раздел, 3 — подраздел


@dataclass
class Para:
    text: str
    bold: str = ""          # жирное начало абзаца
    meta: bool = False      # мелкий серый текст (служебная строка)


@dataclass
class Note:
    """Жёлтая заметка: «Допущения: …», «Подсказка: …». bold — весь текст жирным."""
    title: str
    text: str
    bold: bool = False


@dataclass
class Table:
    head: list[str]
    rows: list[list[str]]                 # уже отформатированные строки
    num_from: int = 1                     # колонки с этого номера — числа (выравнивание вправо)
    num_cols: set[int] | None = None      # или явный набор числовых колонок
    wide: bool = False                    # много колонок — мельче шрифт в Word

    def is_num(self, i: int) -> bool:
        return i in self.num_cols if self.num_cols is not None else i >= self.num_from


@dataclass
class Image:
    png: bytes
    alt: str = ""
    caption: str = ""       # жирная подпись над рисунком
    text: str = ""          # обычный текст после подписи


@dataclass
class Bullets:
    items: list[tuple[str, str]]          # (жирная метка, текст)


@dataclass
class Flags:
    items: list[tuple[str, str, str]]     # (уровень ERROR/WARN/INFO, образцы, текст)
    level_names: dict = field(default_factory=dict)


def fig_png(fig, dpi=110) -> bytes:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=dpi, facecolor="white")
    return buf.getvalue()


# ---------------------------------------------------------------- HTML
LEVEL_CSS = {"ERROR": "err", "WARN": "warn", "INFO": "info"}

CSS = """
body{font-family:Tahoma,Verdana,Arial,sans-serif;font-size:14px;color:#000;background:#fff;
     max-width:1100px;margin:0 auto;padding:16px;line-height:1.45}
h1{font-size:22px;background:#000080;color:#fff;padding:6px 10px;margin:0 0 4px}
h2{font-size:17px;border-bottom:2px solid #000080;padding-bottom:2px;margin-top:28px}
h3{font-size:15px;margin:18px 0 4px}
.meta{color:#444;font-size:13px}
table{border-collapse:collapse;font-size:12.5px;margin:6px 0}
th,td{border:1px solid #808080;padding:3px 6px;vertical-align:top}
th{background:#c0c0c0;text-align:center}
td.n{text-align:right;font-variant-numeric:tabular-nums;white-space:nowrap}
.scroll{overflow-x:auto}
.badge{display:inline-block;min-width:76px;text-align:center;font-weight:bold;padding:0 4px;
       border:1px solid #000;margin-right:6px}
.err{background:#ff0000;color:#fff}.warn{background:#ffff00;color:#000}.info{background:#0000ff;color:#fff}
img{max-width:100%;border:1px solid #808080}
.card{margin:14px 0 24px}
.note{background:#ffffe1;border:1px solid #000;padding:6px 10px;font-size:13px}
ul.flags{list-style:none;padding-left:0}
ul.flags li{margin:3px 0}
@media print{body{max-width:none}h2{page-break-after:avoid}.card,img,table{page-break-inside:avoid}}
"""


def _e(s) -> str:
    return html.escape(str(s))


def _b64(png: bytes) -> str:
    import base64

    return base64.b64encode(png).decode("ascii")


def render_html(blocks, title="Отчёт по гранулометрии") -> str:
    h = [f"<!doctype html><html lang='ru'><head><meta charset='utf-8'>"
         f"<meta name='viewport' content='width=device-width,initial-scale=1'>"
         f"<title>{_e(title)}</title><style>{CSS}</style></head><body>"]
    for b in blocks:
        if isinstance(b, Heading):
            h.append(f"<h{b.level}>{_e(b.text)}</h{b.level}>")
        elif isinstance(b, Para):
            cls = " class='meta'" if b.meta else ""
            bold = f"<b>{_e(b.bold)}</b> " if b.bold else ""
            h.append(f"<p{cls}>{bold}{_e(b.text)}</p>")
        elif isinstance(b, Note):
            if b.bold:
                h.append(f"<p class='note'><b>{_e(b.title)} {_e(b.text)}</b></p>")
            else:
                h.append(f"<p class='note'><b>{_e(b.title)}</b> {_e(b.text)}</p>")
        elif isinstance(b, Table):
            t = ["<div class='scroll'><table><tr>"] + [f"<th>{_e(x)}</th>" for x in b.head] + ["</tr>"]
            for r in b.rows:
                t.append("<tr>" + "".join(f"<td class='n'>{_e(v)}</td>" if b.is_num(i) else f"<td>{_e(v)}</td>"
                                          for i, v in enumerate(r)) + "</tr>")
            t.append("</table></div>")
            h.append("".join(t))
        elif isinstance(b, Image):
            img = f"<img alt='{_e(b.alt)}' src='data:image/png;base64,{_b64(b.png)}'>"
            if b.caption or b.text:
                cap = f"<b>{_e(b.caption)}</b>" + (f"{_e(b.text)}" if b.text else "")
                h.append(f"<div class='card'>{cap}<br>{img}</div>")
            else:
                h.append(img)
        elif isinstance(b, Bullets):
            h.append("<ul>" + "".join(f"<li><b>{_e(k)}</b>: {_e(v)}</li>" for k, v in b.items) + "</ul>")
        elif isinstance(b, Flags):
            h.append("<ul class='flags'>")
            for lv, labels, text in b.items:
                h.append(f"<li><span class='badge {LEVEL_CSS.get(lv, 'info')}'>{_e(b.level_names.get(lv, lv))}"
                         f"</span><b>{_e(labels)}</b>: {_e(text)}</li>")
            h.append("</ul>")
    h.append("</body></html>")
    return "\n".join(h)


def write_html(blocks, path: Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(render_html(blocks), encoding="utf-8")
    return path


# ---------------------------------------------------------------- Word (.docx)
FLAG_FILL = {"ERROR": ("FF0000", "FFFFFF"), "WARN": ("FFFF00", "000000"), "INFO": ("0000FF", "FFFFFF")}


def _shade(el, fill: str):
    """Заливка ячейки или абзаца (w:shd) — python-docx не даёт это сделать без XML."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    pr = el.get_or_add_tcPr() if el.tag.endswith("}tc") else el.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pr.append(shd)


def _box(paragraph):
    """Тонкая чёрная рамка вокруг абзаца (заметки «Допущения:»)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    pbdr = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "4")
        e.set(qn("w:color"), "000000")
        pbdr.append(e)
    paragraph._p.get_or_add_pPr().append(pbdr)


def _repeat_header(row):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    trpr = row._tr.get_or_add_trPr()
    e = OxmlElement("w:tblHeader")
    e.set(qn("w:val"), "true")
    trpr.append(e)


def _col_weights(head, rows) -> list[float]:
    """Относительные ширины колонок: по самому длинному слову заголовка и по содержимому ячеек."""
    w = []
    for i, name in enumerate(head):
        word = 1.15 * max((len(x) for x in str(name).split()), default=1)   # заголовок жирный — шире
        cells = [len(str(r[i])) for r in rows if i < len(r)]
        content = min(max(cells, default=1), 22)
        w.append(max(word, content, 4) + 3)   # + поля ячейки
    return w


def _cant_split(row):
    from docx.oxml import OxmlElement

    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def _no_borders(table):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tblpr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "nil")
        b.append(e)
    tblpr.append(b)


def render_docx(blocks, path: Path, title="Отчёт по гранулометрии") -> Path:
    """Документ Word: А4 альбомная (широкие таблицы), стили «Заголовок 1/2» — по ним Word строит оглавление."""
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    doc.core_properties.title = title
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(1.8))
    sec.top_margin = sec.bottom_margin = Cm(1.5)
    text_w = sec.page_width - sec.left_margin - sec.right_margin

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(4)
    for name, size in (("Title", 18), ("Heading 1", 14), ("Heading 2", 12)):
        st = doc.styles[name]
        st.font.name = "Arial"
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(0, 0, 0x80)

    # номер страницы в нижнем колонтитуле
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            e = OxmlElement("w:fldChar")
            e.set(qn("w:fldCharType"), kind)
        else:
            e = OxmlElement("w:instrText")
            e.set(qn("xml:space"), "preserve")
            e.text = text
        run._r.append(e)

    def card_cell(par, b):
        par.paragraph_format.keep_with_next = True
        if b.caption:
            par.add_run(b.caption).bold = True
        if b.text:
            r = par.add_run(b.text)
            r.font.size = Pt(9)

    i = 0
    while i < len(blocks):
        b = blocks[i]
        if isinstance(b, Image) and (b.caption or b.text):
            # графики образцов — по два в ряд (таблица без рамок), чтобы отчёт не растягивался на десятки страниц
            cards = [b]
            while i + 1 < len(blocks) and isinstance(blocks[i + 1], Image) and blocks[i + 1].caption:
                i += 1
                cards.append(blocks[i])
            for k in range(0, len(cards), 2):
                pair = cards[k:k + 2]
                t = doc.add_table(rows=1, cols=2)
                t.autofit = False
                _no_borders(t)
                _cant_split(t.rows[0])   # подпись и график не разрываются между страницами
                for j in range(2):
                    t.columns[j].width = text_w // 2
                    t.rows[0].cells[j].width = text_w // 2
                for j, c in enumerate(pair):
                    cell = t.rows[0].cells[j]
                    card_cell(cell.paragraphs[0], c)
                    pp = cell.add_paragraph()
                    pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    pp.add_run().add_picture(io.BytesIO(c.png), width=text_w // 2 - Cm(1.6))  # два ряда на страницу
                doc.add_paragraph().paragraph_format.space_after = Pt(0)
            i += 1
            continue
        if isinstance(b, Heading):
            doc.add_heading(b.text, level={1: 0, 2: 1, 3: 2}.get(b.level, 2))
        elif isinstance(b, Para):
            p = doc.add_paragraph()
            if b.bold:
                p.add_run(b.bold + " ").bold = True
            r = p.add_run(b.text)
            if b.meta:
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
        elif isinstance(b, Note):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.15)
            p.paragraph_format.right_indent = Cm(0.15)
            _shade(p._p, "FFFFE1")
            _box(p)
            r = p.add_run(b.title + " ")
            r.bold = True
            r.font.size = Pt(9)
            r = p.add_run(b.text)
            r.bold = b.bold
            r.font.size = Pt(9)
        elif isinstance(b, Table):
            t = doc.add_table(rows=1, cols=len(b.head))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.LEFT
            t.autofit = False
            size = Pt(7 if b.wide else 8.5)
            weights = _col_weights(b.head, b.rows)
            total = sum(weights)
            # узкая таблица (мало колонок) не растягивается на всю ширину страницы
            span = text_w if total > 70 or b.wide else int(text_w * max(0.45, total / 70))
            widths = [int(span * w / total) for w in weights]
            for j, w in enumerate(widths):
                t.columns[j].width = w
            for j, name in enumerate(b.head):
                c = t.rows[0].cells[j]
                c.width = widths[j]
                r = c.paragraphs[0].add_run(name)
                r.bold, r.font.size = True, size
                c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                c.paragraphs[0].paragraph_format.space_after = Pt(0)
                _shade(c._tc, "C0C0C0")
            _repeat_header(t.rows[0])
            for row in b.rows:
                cells = t.add_row().cells
                for j, v in enumerate(row):
                    cells[j].width = widths[j]
                    par = cells[j].paragraphs[0]
                    par.paragraph_format.space_after = Pt(0)
                    r = par.add_run(str(v))
                    r.font.size = size
                    if b.is_num(j):
                        par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
        elif isinstance(b, Image):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(io.BytesIO(b.png), height=Cm(11))
        elif isinstance(b, Bullets):
            for k, v in b.items:
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.space_after = Pt(1)
                p.add_run(k).bold = True
                p.add_run(": " + v)
        elif isinstance(b, Flags):
            for lv, labels, text in b.items:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(2)
                fill, color = FLAG_FILL.get(lv, FLAG_FILL["INFO"])
                r = p.add_run(f" {b.level_names.get(lv, lv)} ")
                r.bold = True
                r.font.color.rgb = RGBColor.from_string(color)
                rpr = r._r.get_or_add_rPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:color"), "auto")
                shd.set(qn("w:fill"), fill)
                rpr.append(shd)
                p.add_run("  ")
                p.add_run(labels).bold = True
                p.add_run(": " + text)
        i += 1
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path
