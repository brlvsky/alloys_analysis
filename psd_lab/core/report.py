"""Отчёты: summary.xlsx (Сводка / Кривые / Флаги), самодостаточный report.html и report.docx (Word).

Плюс пакетный режим: обработать папку целиком без GUI.
"""
from __future__ import annotations

import datetime as dt
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np

from .. import APP_NAME, PROJECT, __version__
from . import expand_paths, load
from .metrics import average_repeats, compute, window_label
from .model import LEVEL_NAMES, LEVEL_ORDER, Sample
from .plots import common_axes, draw_compare, draw_sample, new_figure, plot_all

SOURCE_NAMES = {"fritsch": "Fritsch", "table": "таблица"}


# ---------------------------------------------------------------- данные сводки
def summary_columns(windows, with_qc=False) -> list[str]:
    cols = ["Образец", "Измерения", "Файл", "d10, мкм", "d50, мкм", "d90, мкм", "span",
            "D[4,3], мкм", "D[3,2], мкм"]
    cols += [f"{window_label(lo, hi)} мкм, %" for lo, hi in windows]
    cols += ["Обскурация, %", "Error", "Флаги", "Источник"]
    if with_qc:
        cols.append("QC")
    return cols


def flags_short(s: Sample) -> str:
    """Кратко: 'ОШИБКА 1, ИНФО 1' или пусто."""
    if not s.flags:
        return ""
    counts = {}
    for f in s.flags:
        counts[f[0]] = counts.get(f[0], 0) + 1
    order = sorted(counts, key=lambda lv: LEVEL_ORDER.get(lv, 9))
    return ", ".join(f"{LEVEL_NAMES.get(lv, lv)} {counts[lv]}" for lv in order)


def summary_rows(samples: list[Sample], windows, qc: dict | None = None) -> list[list]:
    """qc — {подпись образца: «3/5»}; если задан, добавляется колонка QC."""
    rows = []
    for s in samples:
        m = compute(s, windows)
        obs = s.meta.get("obscuration")
        err = s.meta.get("error")
        rows.append([
            s.name, ", ".join(s.members) if s.members else s.meas_id, Path(s.file).name if s.file else "",
            m["d10"], m["d50"], m["d90"], m["span"], m["d43"], m["d32"],
            *m["fractions"].values(),
            obs, err, flags_short(s), SOURCE_NAMES.get(s.source, s.source),
        ] + ([qc.get(s.label, "")] if qc is not None else []))
    return rows


def grouped_flags(samples: list[Sample]) -> list[tuple[tuple[str, str], list[str]]]:
    """Одинаковые флаги разных образцов (например, TradeOff файла) — одной строкой."""
    groups: dict[tuple[str, str], list[str]] = {}
    for s in samples:
        for f in s.flags:
            groups.setdefault((f[0], f[1]), []).append(s.label)
    return sorted(groups.items(), key=lambda kv: (LEVEL_ORDER.get(kv[0][0], 9), kv[1][0]))


def interval_q(s: Sample) -> np.ndarray:
    """Q, % в интервале, заканчивающемся в каждой точке сетки (первая точка — 0)."""
    return np.concatenate([[0.0], np.diff(s.cum_pct)])


# ---------------------------------------------------------------- Excel
def write_xlsx(samples: list[Sample], path: Path, windows, st=None, qc=None) -> Path:
    """Сводка, кривые, флаги; с настройками st — ещё листы модулей М2–М5."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head_font = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="D9D9D9")
    wb = Workbook()

    ws = wb.active
    ws.title = "Сводка"
    cols = summary_columns(windows, with_qc=qc is not None)
    ws.append(cols)
    for r in summary_rows(samples, windows, qc):
        ws.append([_xl(v) for v in r])
    for c in range(1, len(cols) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font, cell.fill = head_font, head_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(c)].width = 22 if c in (1, 2, 3) else 11
    ws.column_dimensions[get_column_letter(cols.index("Флаги") + 1)].width = 24
    ws.column_dimensions[get_column_letter(cols.index("Источник") + 1)].width = 24
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=len(cols) - 2):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "0.00"
    ws.freeze_panes = "B2"

    wc = wb.create_sheet("Кривые")
    for i, s in enumerate(samples):
        c0 = i * 4 + 1
        wc.cell(row=1, column=c0, value=s.label).font = head_font
        for j, h in enumerate(("Размер, мкм", "ΣQ, %", "Q, %")):
            cell = wc.cell(row=2, column=c0 + j, value=h)
            cell.font, cell.fill = head_font, head_fill
        q = interval_q(s)
        for k, (x, c, qq) in enumerate(zip(s.size_um, s.cum_pct, q)):
            wc.cell(row=3 + k, column=c0, value=float(x))
            wc.cell(row=3 + k, column=c0 + 1, value=round(float(c), 4))
            wc.cell(row=3 + k, column=c0 + 2, value=round(float(qq), 4))
    wc.freeze_panes = "A3"

    wf = wb.create_sheet("Флаги")
    wf.append(["Образец", "Измерения", "Уровень", "Описание"])
    for c in range(1, 5):
        wf.cell(row=1, column=c).font, wf.cell(row=1, column=c).fill = head_font, head_fill
    for s in samples:
        for f in sorted(s.flags, key=lambda f: LEVEL_ORDER.get(f[0], 9)):
            wf.append([s.name, ", ".join(s.members) or s.meas_id, LEVEL_NAMES.get(f[0], f[0]), f[1]])
    for col, w in zip("ABCD", (22, 20, 12, 100)):
        wf.column_dimensions[col].width = w

    if st is not None and samples:
        from .report_modules import xlsx_sheets

        xlsx_sheets(wb, samples, st, head_font, head_fill)

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _xl(v):
    if isinstance(v, (float, np.floating)):
        return None if not np.isfinite(v) else float(v)
    return v


# ---------------------------------------------------------------- отчёт (HTML и Word)
def _fmt(v, nd=2) -> str:
    if v is None:
        return "—"
    if isinstance(v, (float, np.floating)):
        return "—" if not np.isfinite(v) else f"{v:.{nd}f}".replace(".", ",")
    return str(v)


LEGEND = ("ΣQ — накопленная объёмная доля частиц мельче данного размера, %; Q — доля объёма в интервале, %. "
          "d10, d50, d90 — размеры, мельче которых 10, 50 и 90 % объёма порошка; span = (d90 − d10)/d50 — ширина "
          "распределения. D[4,3] — средний по объёму диаметр, D[3,2] — средний по поверхности (Заутера). Доли по "
          "окнам — объёмный процент частиц в диапазоне размеров. Кривые не обрезаются и не перенормируются: если "
          "кривая не доходит до 100 %, это показано флагом. Повторные измерения с одинаковым названием усреднены.")


def build_blocks(samples: list[Sample], *, windows, lang="en", bin_um=None, xmax=None, independent_axes=False,
                 show_name=True, log_x=True, files=None, skipped=None, st=None, batches=None, qc=None,
                 dist_log=None) -> list:
    """Содержание отчёта — список блоков (report_doc). Из него делаются и HTML, и Word."""
    from .report_doc import Bullets, Flags, Heading, Image, Note, Para, Table, fig_png

    files = files or sorted({s.file for s in samples})
    now = dt.datetime.now().strftime("%d.%m.%Y %H:%M")
    cols = summary_columns(windows, with_qc=qc is not None)
    rows = summary_rows(samples, windows, qc)
    nw = len(windows)
    num_cols = set(range(3, 11 + nw))
    decimals = {i: 2 for i in num_cols}
    decimals.update({i: 1 for i in range(9, 9 + nw)})
    decimals[9 + nw] = 0   # обскурация
    decimals[10 + nw] = 3  # Error

    b = [Heading("Гранулометрический состав порошков", 1),
         Para(f"Проект: «{PROJECT}». Сформировано {now} программой {APP_NAME} {__version__}. "
              f"Образцов: {len(samples)}.", meta=True),
         Para("Файлы: " + ", ".join(Path(f).name for f in files), meta=True)]

    # Сводка
    b.append(Heading("1. Сводная таблица"))
    b.append(Table(cols, [[_fmt(v, decimals[i]) if i in num_cols else _fmt(v) for i, v in enumerate(r)]
                          for r in rows], num_cols=num_cols, wide=True))
    b.append(Note("Обозначения.", LEGEND))

    # Флаги
    b.append(Heading("2. Качество измерений"))
    if not any(s.flags for s in samples):
        b.append(Para("Замечаний нет."))
    else:
        b.append(Flags([(lv, ", ".join(labels), text) for (lv, text), labels in grouped_flags(samples)],
                       level_names=LEVEL_NAMES))
        b.append(Note("Уровни:", "ОШИБКА — результату доверять нельзя, измерение нужно повторить. ВНИМАНИЕ — результат "
                      "возможен, но с оговоркой. ИНФО — к сведению при сравнении образцов. Поле Error прибора "
                      "показано в таблице без оценки (единицы в экспорте не указаны)."))

    # Сравнение
    groups = _groups(samples)
    fig = new_figure()
    draw_compare(fig, groups, lang=lang, log_x=log_x, xmax=xmax)
    b.append(Heading("3. Сравнение накопленных кривых"))
    b.append(Image(fig_png(fig), alt="Сравнение"))
    pair = report_pair(samples, st)
    if pair is not None:
        from . import compare2
        from .plots import draw_pair

        d = compare2.compare(*pair, windows)
        b.append(Heading(f"До и после: A — {d.a.label}, B — {d.b.label}", 3))
        draw_pair(fig, d, lang=lang, log_x=log_x)
        b.append(Image(fig_png(fig), alt="До и после"))
        b.append(Table(["Показатель", "A", "B", "B − A", "Изменение, %"], compare2.table_rows(d)))
        b.append(Note("Итог:", compare2.summary_text(d)))
        for w in d.warnings:
            b.append(Note("Внимание:", w + "."))
        b.append(Note("Допущения:", compare2.ASSUMPTIONS))

    # По образцам
    b.append(Heading("4. Распределения по образцам"))
    for grp in groups:
        if dist_log is None:
            dist_log = bool(getattr(st, "dist_log", False))
        shared = (None, None, None) if independent_axes else common_axes(grp, bin_um, xmax, dist_log)
        b.append(Heading(f"Файл {Path(grp[0].file).name}", 3))
        for s in grp:
            draw_sample(fig, s, lang=lang, bin_um=bin_um, xmax=xmax or shared[0], ymax=shared[1],
                        show_name=show_name, log_x=dist_log, xmin=shared[2])
            m = compute(s, windows)
            text = (f" — d10 {_fmt(m['d10'])}, d50 {_fmt(m['d50'])}, d90 {_fmt(m['d90'])} мкм; "
                    f"D[4,3] {_fmt(m['d43'])} мкм")
            if s.meta.get("obscuration") is not None:
                text += f"; обскурация {_fmt(s.meta['obscuration'], 0)} %"
            b.append(Image(fig_png(fig), alt=s.label, caption=s.label, text=text))

    n = 5
    if st is not None and samples:
        from .report_modules import module_blocks

        more, n = module_blocks(samples, st, fig, start=5)
        b += more
    if batches:
        from .report_modules import batches_blocks

        b += batches_blocks(batches, n)

    if skipped:
        b.append(Heading("Пропущенные файлы"))
        b.append(Bullets([(Path(f).name, why) for f, why in skipped]))
    return b


def report_pair(samples: list[Sample], st) -> tuple[Sample, Sample] | None:
    """Пара «до и после» для отчёта — если она выбрана во вкладке «Сравнение» и оба образца в отчёте."""
    pair = list(getattr(st, "compare_pair", None) or [])
    by = {s.label: s for s in samples}
    if len(pair) == 2 and pair[0] in by and pair[1] in by and pair[0] != pair[1]:
        return by[pair[0]], by[pair[1]]
    return None


def write_reports(samples: list[Sample], html_path: Path | None = None, docx_path: Path | None = None,
                  **kw) -> list[Path]:
    """HTML и Word из одного набора блоков (графики рисуются один раз)."""
    from .report_doc import render_docx
    from .report_doc import write_html as _write

    blocks = build_blocks(samples, **kw)
    out = []
    if html_path:
        out.append(_write(blocks, html_path))
    if docx_path:
        out.append(render_docx(blocks, docx_path))
    return out


def write_html(samples: list[Sample], path: Path, **kw) -> Path:
    """Самодостаточный report.html (картинки внутри файла)."""
    from .report_doc import write_html as _write

    return _write(build_blocks(samples, **kw), path)


def write_docx(samples: list[Sample], path: Path, **kw) -> Path:
    """report.docx — тот же отчёт в Word (можно править и сохранить в PDF средствами Word)."""
    from .report_doc import render_docx

    return render_docx(build_blocks(samples, **kw), path)


def _groups(samples: list[Sample]) -> list[list[Sample]]:
    groups: dict[str, list[Sample]] = {}
    for s in samples:
        groups.setdefault(s.file, []).append(s)
    return list(groups.values())


# ---------------------------------------------------------------- пакетный режим
@dataclass
class BatchResult:
    samples: list = field(default_factory=list)
    files: list = field(default_factory=list)
    skipped: list = field(default_factory=list)   # [(файл, причина)]
    outputs: list = field(default_factory=list)


def load_many(paths, average=True, log=print, recipes=None) -> BatchResult:
    """recipes — {SHA-1 файла: ручная настройка импорта} (из настроек программы)."""
    from . import file_sha1

    res = BatchResult()
    for f in expand_paths(paths):
        try:
            rec = (recipes or {}).get(file_sha1(f)) if recipes else None
            ss = load(f, rec)
        except Exception as e:  # noqa: BLE001 — файл пользователя может быть любым
            res.skipped.append((f, f"не удалось прочитать: {e}"))
            log(f"ПРОПУЩЕН  {f.name}: не удалось прочитать ({e})")
            continue
        if not ss:
            res.skipped.append((f, "не найдено распределений"))
            log(f"ПРОПУЩЕН  {f.name}: не найдено распределений")
            continue
        src = "экспорт Fritsch" if ss[0].source == "fritsch" else ("ручная настройка" if rec else "таблица")
        log(f"ЗАГРУЖЕН  {f.name}: {len(ss)} изм. ({src})")
        for note in dict.fromkeys(n for s in ss for n in s.meta.get("import_notes", [])):
            log(f"          ↳ {note}")
        res.files.append(f)
        res.samples.extend(average_repeats(ss) if average else ss)
    return res


def run_batch(src, out, settings=None, log=print, db_file=None) -> BatchResult:
    from .settings import Settings

    st = settings or Settings()
    out = Path(out)
    res = load_many([src] if isinstance(src, (str, Path)) else src, average=st.average, log=log,
                    recipes=st.import_recipes)
    if not res.samples:
        log("Нет данных для обработки.")
        return res
    batches = None
    if db_file:
        from . import db, file_sha1

        conn = db.connect(db_file)
        for f in res.files:
            n = db.import_samples(conn, load(f), file_sha1(f), st)
            log(f"БАЗА      {Path(f).name}: добавлено измерений: {n}" if n else f"БАЗА      {Path(f).name}: уже в базе")
        names = list(dict.fromkeys(s.name for s in res.samples))
        batches = [dict(r) for n in names for r in conn.execute("SELECT * FROM batches WHERE name=?", (n,))]
        conn.close()
    groups = [(g[0].file, g) for g in _groups(res.samples)]
    made = plot_all(groups, out, lang=st.lang, bin_um=st.bin_um, xmax=st.xmax,
                    independent_axes=st.independent_axes, show_name=st.show_name, log_x=st.compare_log,
                    dist_log=st.dist_log)
    res.outputs += [p for _, p in made]
    log(f"Графики: {len(made)} PNG в {out}")
    log("Модули: популяции, окна печати, выход годного, поверхность…")
    res.outputs.append(write_xlsx(res.samples, out / "summary.xlsx", st.windows_tuples, st=st))
    res.outputs += write_reports(res.samples, out / "report.html", out / "report.docx", windows=st.windows_tuples,
                                 lang=st.lang, bin_um=st.bin_um, xmax=st.xmax, independent_axes=st.independent_axes,
                                 show_name=st.show_name, log_x=st.compare_log, files=res.files,
                                 skipped=res.skipped, st=st, batches=batches)
    log(f"Отчёт: {out / 'report.html'} и {out / 'report.docx'}; таблица: {out / 'summary.xlsx'}")
    return res
