"""Вывод расчёта шихты: бланк навесок (Excel, Word, HTML) и книга Excel со всеми вариантами.

Блоки документа — те же, что у отчёта по гранулометрии (`report_doc`), поэтому бланк одинаково
выглядит в Word, HTML и при печати.
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path

from .. import APP_NAME, __version__
from . import charge
from .charge import Result, balance_accuracy, fmt_g, fmt_pct

HEAD_FILL = "C0C0C0"


# ==================================================================== бланк навесок (блоки документа)
def blank_blocks(res: Result, per_load: bool = True, variant=None) -> list:
    """Бланк навесок: шапка, таблица «компонент | масса | точность весов | фактически | подпись».
    per_load — отдельный лист на каждую загрузку, иначе одна таблица на всю партию."""
    from .report_doc import Heading, Note, Para, Table

    v = variant or res.main
    r, plan = res.recipe, res.plan
    now = dt.datetime.now().strftime("%d.%m.%Y")
    out = [Heading("Бланк навесок шихты", 1)]
    out.append(Para(f"Сплав: {r.alloy or '—'}. Вариант: {v.name}. Дата: {now}. "
                    f"Программа {APP_NAME} {__version__}.", meta=True))
    rows_of = []
    if per_load:
        for k in range(1, plan.n + 1):
            rows_of.append((f"Загрузка {k} из {plan.n} — {fmt_g(plan.per_load_g)} г", res.loads(v)))
    else:
        rows_of.append((f"Вся партия: {plan.n} загрузок по {fmt_g(plan.per_load_g)} г = "
                        f"{fmt_g(plan.made_g)} г", res.purchase(v)))
    for title, masses_g in rows_of:
        out.append(Heading(title, 3))
        rows = [[c, fmt_g(g), balance_accuracy(g), "", ""] for c, g in masses_g.items()]
        rows.append(["Итого", fmt_g(sum(masses_g.values())), "", "", ""])
        out.append(Table(["Компонент", "Навеска, г", "Точность весов", "Фактически взвешено, г", "Подпись"],
                         rows, num_from=1))
    if plan.balls_g:
        out.append(Para(f"Шары: {fmt_g(plan.balls_g)} г на загрузку "
                        f"(соотношение шары : порошок = {plan.balls_ratio:g} : 1)."))
    if plan.note:
        out.append(Note("Внимание:", plan.note))
    for lv, text in v.flags:
        out.append(Note("Внимание:" if lv == "WARN" else "Ошибка:", text))
    for w in res.warnings:
        out.append(Note("Проверьте ввод:", w))
    out.append(Note("Допущения:", charge.ASSUMPTIONS))
    return out


def write_blank_docx(res: Result, path: Path, per_load: bool = True, variant=None) -> Path:
    from .report_doc import render_docx

    return render_docx(blank_blocks(res, per_load, variant), path, title="Бланк навесок шихты")


def write_blank_html(res: Result, path: Path, per_load: bool = True, variant=None) -> Path:
    from .report_doc import write_html

    return write_html(blank_blocks(res, per_load, variant), path)


# ==================================================================== Excel
def write_xlsx(res: Result, path: Path, per_load: bool = True, variant=None) -> Path:
    """Книга: «Состав», «Варианты», «Загрузки», «Закупка», «Лигатура», «Бланк навесок».
    variant — по какому варианту считать навески (по умолчанию последний, как на вкладке)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head_font = Font(bold=True)
    head_fill = PatternFill("solid", fgColor=HEAD_FILL)
    wb = Workbook()
    wb.remove(wb.active)
    r, plan = res.recipe, res.plan
    base, main = res.variants[0], variant or res.main

    def sheet(title, head, rows, widths=None):
        ws = wb.create_sheet(title[:31])
        ws.append(head)
        for c in range(1, len(head) + 1):
            ws.cell(row=1, column=c).font = head_font
            ws.cell(row=1, column=c).fill = head_fill
            ws.cell(row=1, column=c).alignment = Alignment(horizontal="center", wrap_text=True)
        for row in rows:
            ws.append(row)
        for i, w in enumerate(widths or [], start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        return ws

    # ---- Состав
    comps = list(base.w)
    ws = sheet("Состав", ["Компонент", "ат.% / мол.%", "мас.%", "На 200 г, г", "Чистота, %"],
               [[c, round(base.x[c], 6), round(base.w[c], 6), round(base.w[c] * 2, 4),
                 (r.purity or {}).get(c, 100)] for c in comps], [16, 14, 14, 14, 12])
    ws.append(["Итого", round(sum(base.x.values()), 6), round(sum(base.w.values()), 6), 200, None])
    ws.append([])
    ws.append([f"Сплав: {r.alloy}"])
    ws.append([f"Состав: {r.composition} ({charge.UNITS[r.basis]})"])
    ws.append([f"Σ xᵢ·Mᵢ = {charge.sum_xm(base.x, charge.masses(r.mass_overrides)):.4f}"])
    if r.mass_overrides:
        ws.append(["Переопределённые атомные массы: " +
                   ", ".join(f"{k} {v:g}" for k, v in r.mass_overrides.items())])
    for w in res.warnings:
        ws.append([w])

    # ---- Варианты
    all_comps = list(dict.fromkeys(c for v in res.variants for c in v.w))
    head = ["Вариант"] + [f"{c}, мас.%" for c in all_comps] + ["O из оксидов, мас.%", "ρ, г/см³", "Δ к Ti6Al4V, %"]
    rows = []
    for v in res.variants:
        rho = v.density.rho if v.density else None
        rows.append([v.name] + [round(v.w[c], 6) if c in v.w else None for c in all_comps]
                    + [round(v.oxygen_total, 6) if v.oxygen else None,
                       round(rho, 4) if rho else None,
                       round(v.density.delta_pct, 2) if v.density else None])
    ws = sheet("Варианты", head, rows, [34] + [12] * (len(head) - 1))
    ws.append([])
    for v in res.variants:
        for lv, text in v.flags:
            ws.append([f"{v.name}: {text}"])

    # ---- Загрузки
    loads = res.loads(main)
    ws = sheet("Загрузки", ["Компонент", "На загрузку, г", "Точность весов"],
               [[c, round(g, 4), balance_accuracy(g)] for c, g in loads.items()], [16, 16, 22])
    ws.append(["Итого", round(sum(loads.values()), 4), None])
    ws.append([])
    for line in (f"Вариант: {main.name}",
                 f"Нужно смеси: {fmt_g(plan.target_g)} г; ёмкость барабана: {fmt_g(plan.capacity_g)} г; "
                 f"запас: {plan.reserve_pct:g} %",
                 f"Режим: {'полные загрузки' if plan.mode == 'full' else 'равные загрузки'}; "
                 f"загрузок: {plan.n} по {fmt_g(plan.per_load_g)} г; получится {fmt_g(plan.made_g)} г; "
                 f"избыток {fmt_g(plan.excess_g)} г",
                 (f"Шары: {fmt_g(plan.balls_g)} г на загрузку" if plan.balls_g else ""), plan.note):
        if line:
            ws.append([line])

    # ---- Закупка
    buy = res.purchase(main)
    ws = sheet("Закупка", ["Компонент", f"На {plan.n} загрузок, г"],
               [[c, round(g, 4)] for c, g in buy.items()], [16, 20])
    ws.append(["Итого", round(sum(buy.values()), 4)])

    # ---- Лигатура
    if res.ligature is not None:
        lg = res.ligature
        rows = [["Лигатура", round(lg.lig_pct, 6), round(plan.per_load_g * lg.lig_pct / 100, 4)]]
        rows += [[f"{c} (чистый)", round(v, 6), round(plan.per_load_g * v / 100, 4)] for c, v in lg.pure.items()]
        ws = sheet("Лигатура", ["Что взять", "мас.% шихты", "На загрузку, г"], rows, [20, 16, 16])
        ws.append([])
        ws.append(["Состав лигатуры", "мас.%"])
        for c, v in lg.lig_comp.items():
            ws.append([c, round(v, 6)])
        if any(abs(d) > 1e-9 for d in lg.deviation.values()):
            ws.append([])
            ws.append(["Отклонение получившегося состава от целевого, п.п."])
            for c, d in lg.deviation.items():
                ws.append([c, round(d, 6)])
        for _, text in lg.flags:
            ws.append([text])

    # ---- Бланк навесок
    ws = wb.create_sheet("Бланк навесок")
    ws.append([f"Бланк навесок — {r.alloy or 'шихта'}"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([f"Вариант: {main.name}. Дата: {dt.datetime.now():%d.%m.%Y}"])
    groups = ([(f"Загрузка {k} из {plan.n} — {fmt_g(plan.per_load_g)} г", loads) for k in range(1, plan.n + 1)]
              if per_load else [(f"Вся партия ({plan.n} × {fmt_g(plan.per_load_g)} г)", buy)])
    for title, masses_g in groups:
        ws.append([])
        ws.append([title])
        ws.cell(row=ws.max_row, column=1).font = head_font
        ws.append(["Компонент", "Навеска, г", "Точность весов", "Фактически взвешено, г", "Подпись"])
        for c in range(1, 6):
            ws.cell(row=ws.max_row, column=c).font = head_font
            ws.cell(row=ws.max_row, column=c).fill = head_fill
        for c, g in masses_g.items():
            ws.append([c, round(g, 4), balance_accuracy(g), None, None])
        ws.append(["Итого", round(sum(masses_g.values()), 4), None, None, None])
    for i, w in enumerate([18, 14, 22, 24, 16], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.append([])
    ws.append([f"Допущения: {charge.ASSUMPTIONS}"])

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))
    return path


# ==================================================================== текстовая сводка (для вкладки)
def summary_lines(res: Result, variant=None) -> list[str]:
    """Короткий итог расчёта для панели «Результат»."""
    plan, main = res.plan, variant or res.main
    out = [f"Вариант «{main.name}»: " + "; ".join(f"{c} {fmt_pct(w)} мас.%" for c, w in main.w.items())]
    out.append(f"Загрузок: {plan.n} по {fmt_g(plan.per_load_g)} г — всего {fmt_g(plan.made_g)} г "
               f"(нужно {fmt_g(plan.target_g)} г, избыток {fmt_g(plan.excess_g)} г)")
    if main.density:
        out.append(f"Плотность по правилу смесей: {main.density.rho:.3f} г/см³ "
                   f"({main.density.delta_pct:+.1f} % к Ti6Al4V)".replace(".", ",").replace("-", "−"))
    elif main.density_note:
        out.append(main.density_note)
    return out
