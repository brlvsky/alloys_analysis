"""Чтение файлов-таблиц (.xlsx, .xls, .csv, .txt) в «сетку» ячеек; поиск распределений — io_smart."""
from __future__ import annotations

import csv
import datetime as dt
import io
import re
from pathlib import Path

import numpy as np

from .model import Sample

# ---------------------------------------------------------------- чтение файлов
def read_grids(path: Path) -> list[tuple[str, list[list]]]:
    """Возвращает [(имя листа, строки)], ячейки — float / str / datetime / None."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".xls":
        return _read_xls(path)
    if ext in (".xlsx", ".xlsm"):
        return _read_xlsx(path)
    if ext in (".csv", ".txt", ".tsv", ".dat"):
        return [(path.stem, _read_text(path))]
    raise ValueError(f"Неизвестный формат файла: {path.name}")


def _read_xls(path: Path):
    import xlrd

    book = xlrd.open_workbook(str(path))
    out = []
    for sh in book.sheets():
        rows = []
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    row.append(None)
                elif cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                elif cell.ctype in (xlrd.XL_CELL_NUMBER, xlrd.XL_CELL_BOOLEAN):
                    row.append(float(cell.value))
                elif cell.ctype == xlrd.XL_CELL_TEXT:
                    row.append(cell.value if cell.value.strip() else None)
                else:
                    row.append(None)
            rows.append(row)
        out.append((sh.name, rows))
    return out


def _read_xlsx(path: Path):
    import openpyxl

    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = []
            for v in row:
                if isinstance(v, bool):
                    v = float(v)
                elif isinstance(v, (int, float)):
                    v = float(v)
                elif isinstance(v, str) and not v.strip():
                    v = None
                cells.append(v)
            rows.append(cells)
        out.append((ws.title, rows))
    wb.close()
    return out


def _read_text(path: Path) -> list[list]:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1251", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    lines = [ln for ln in text.splitlines() if ln.strip()]
    sample = "\n".join(lines[:50])
    if "\t" in sample:
        delim = "\t"
    elif ";" in sample:
        delim = ";"
    elif "," in sample and not _whitespace_table(lines[:50]):
        delim = ","  # запятая-разделитель, десятичная точка
    else:
        delim = None  # пробелы (числа могут быть с десятичной запятой)
    decimal_comma = delim != ","
    rows = []
    if delim is None:
        parsed = [ln.split() for ln in lines]
    else:
        parsed = list(csv.reader(io.StringIO("\n".join(lines)), delimiter=delim))
    for parts in parsed:
        rows.append([_parse_cell(p, decimal_comma) for p in parts])
    return rows


def _whitespace_table(lines) -> bool:
    """Строки вида '1,5  2,3' — столбцы через пробелы, десятичная запятая."""
    data = [ln.split() for ln in lines if re.match(r"\s*[\d.,]", ln)]
    return bool(data) and all(len(t) >= 2 and all(re.fullmatch(r"-?\d+(,\d+)?", x) for x in t) for t in data)


def _parse_cell(s: str, decimal_comma: bool):
    s = s.strip().strip('"')
    if not s:
        return None
    t = s.replace(" ", "").replace(" ", "")
    if decimal_comma:
        t = t.replace(",", ".")
    try:
        return float(t)
    except ValueError:
        return s


def _num(v):
    return isinstance(v, float) and np.isfinite(v)


# ---------------------------------------------------------------- разбор таблиц
def parse_grid(rows: list[list], sheet: str, file: Path) -> list[Sample]:
    """Распределения на листе — см. io_smart (оси размеров, виды данных, единицы, таблицы «боком»)."""
    from .io_smart import parse_sheet

    return parse_sheet(rows, sheet, file)


def load_table(path: Path) -> list[Sample]:
    path = Path(path)
    samples = []
    for sheet, rows in read_grids(path):
        samples.extend(parse_grid(rows, sheet, path))
    return samples


def excel_serial_to_datetime(v) -> dt.datetime | None:
    if isinstance(v, dt.datetime):
        return v
    if _num(v):
        return dt.datetime(1899, 12, 30) + dt.timedelta(days=v)
    return None
