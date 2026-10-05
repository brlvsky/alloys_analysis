"""Разделы отчёта и листы Excel для модулей М2–М5."""
from __future__ import annotations

import html

from . import deconv, surface, windows
from .metrics import d32, window_label


def _c(v, nd=1):
    return "—" if v is None or v != v else f"{v:.{nd}f}".replace(".", ",")


def _table(head, rows, num_from=1):
    h = ["<div class='scroll'><table><tr>"] + [f"<th>{html.escape(x)}</th>" for x in head] + ["</tr>"]
    for r in rows:
        h.append("<tr>" + "".join(
            f"<td class='n'>{html.escape(str(v))}</td>" if i >= num_from else f"<td>{html.escape(str(v))}</td>"
            for i, v in enumerate(r)) + "</tr>")
    h.append("</table></div>")
    return "".join(h)


def _note(title, text):
    return f"<p class='note'><b>{html.escape(title)}</b> {html.escape(text)}</p>"


# ==================================================================== HTML
def html_sections(samples, st, fig, fig_b64, start=5) -> list[str]:
    from .plots import draw_populations, draw_tech_bars

    out = []
    n = start
    # ---------- М2
    out.append(f"<h2>{n}. Популяции частиц (разложение на логнормальные компоненты)</h2>")
    rows, multi = [], []
    for s in samples:
        r = deconv.fit(s)
        fine, coarse = r.fine(), r.coarse()
        rows.append([s.label, r.k, _c(r.r2, 4), len(r.main_populations),
                     _c(fine.weight_pct) if fine else "—", _c(fine.mode_um, 1) if fine else "—",
                     _c(coarse.weight_pct) if coarse else "—", _c(coarse.mode_um, 0) if coarse else "—"])
        if r.bimodal:
            multi.append((s, r))
    out.append(_table(["Образец", "Компонент (BIC)", "R²", "Популяций", "Мелкая, %", "мода, мкм",
                       "Крупная, %", "мода, мкм"], rows))
    out.append(_note("Подсказка.", "Доля мелкой популяции у П/С (30–80 %) на порядок больше объёмной доли добавки "
                     "(0,5–1,5 мас.% ≈ 1–3 об.%), значит мелкая мода — не сама добавка; её происхождение "
                     "(агломераты? продукт обработки?) надо проверять СЭМ и измерением с разной мощностью "
                     "ультразвука."))
    out.append(_note("Допущения:", "каждая популяция описывается логнормальным законом; число компонент — по BIC; "
                     "популяции — группы компонент между провалами плотности q3*. Компоненты < 3 % отдельно не "
                     "показываются."))
    for s, r in multi:
        draw_populations(fig, s, r, lang=st.lang)
        out.append(f"<div class='card'><b>{html.escape(s.label)}</b><br>"
                   f"<img alt='' src='data:image/png;base64,{fig_b64(fig)}'></div>")
    n += 1

    # ---------- М3
    sls, ebm = st.sls, st.ebm
    out.append(f"<h2>{n}. Технология: окна СЛС и СЭЛС</h2>")
    draw_tech_bars(fig, samples, sls, ebm)
    out.append(f"<img alt='Окна печати' src='data:image/png;base64,{fig_b64(fig)}'>")
    wins = [("СЛС", tuple(w)) for w in st.sls_windows] + [("СЭЛС", tuple(w)) for w in st.ebm_windows]
    head = ["Образец"] + [f"{t} {w[0]:g}–{w[1]:g}, %" for t, w in wins]
    out.append(_table(head, [[s.label] + [_c(windows.frac(s, *w)) for _, w in wins] for s in samples]))
    out.append("<ul>" + "".join(f"<li><b>{html.escape(s.label)}</b>: {html.escape(windows.facts_text(s, sls, ebm))}"
                                "</li>" for s in samples) + "</ul>")
    out.append(_note("Справка (проверить по источникам):", windows.TECH_NOTE.split(": ", 1)[1]))
    n += 1

    # ---------- М4
    lo, hi = st.sieve_window
    out.append(f"<h2>{n}. Выход годного после рассева {lo:g}–{hi:g} мкм</h2>")
    rows = []
    for s in samples:
        sv = windows.sieve(s, lo, hi)
        after = windows.requirements_check(sv.sieved, st.requirements)
        before = windows.requirements_check(s, st.requirements)
        rows.append([s.label, _c(sv.yield_pct), _c(sv.fines_pct), _c(sv.coarse_pct), f"{sv.grams_per_kg():.0f}",
                     f"{sum(r.ok for r in before)}/{len(before)}", f"{sum(r.ok for r in after)}/{len(after)}"])
    out.append(_table(["Образец", "Годное, %", "Мелочь, %", "Крупное, %", "Из 1 кг, г",
                       "Требования: исходный", "после рассева"], rows))
    req = st.requirements
    out.append(_note("Требования:", f"d10 ≥ {req['d10_min']:g} мкм; d50 {req['d50_min']:g}–{req['d50_max']:g} мкм; "
                     f"d90 ≤ {req['d90_max']:g} мкм; мельче {req['d10_min']:g} мкм ≤ 10 %. " + windows.REQ_NOTE))
    out.append(_note("Допущения:", windows.SIEVE_NOTE))
    n += 1

    # ---------- М5
    rho = st.density_g_cm3
    out.append(f"<h2>{n}. Удельная поверхность и риск по кислороду</h2>")
    rows = []
    for s in samples:
        v, a = surface.fine_shares(s)
        rows.append([s.label, _c(d32(s), 2), _c(surface.ssa_m2_g(s, rho), 3), _c(v), _c(a)])
    out.append(_table(["Образец", "D[3,2], мкм", f"SSA при ρ = {_c(rho, 2)} г/см³, м²/г", "< 15 мкм: % объёма",
                       "< 15 мкм: % поверхности"], rows))
    out.append("<ul>" + "".join(f"<li><b>{html.escape(s.label)}</b>: {html.escape(surface.headline(s))}</li>"
                                for s in samples) + "</ul>")
    out.append(_note("Формула:", "SSA [м²/г] = 6 / (ρ [г/см³] · D[3,2] [мкм]) — нижняя оценка для сферических частиц."))
    out.append(_note("Допущения:", surface.ASSUMPTIONS.split(": ", 1)[1]))
    return out


# ==================================================================== Excel
def xlsx_sheets(wb, samples, st, head_font, head_fill) -> None:
    def sheet(title, head, rows, widths=None):
        ws = wb.create_sheet(title)
        ws.append(head)
        for c in range(1, len(head) + 1):
            ws.cell(row=1, column=c).font, ws.cell(row=1, column=c).fill = head_font, head_fill
        for r in rows:
            ws.append(r)
        from openpyxl.utils import get_column_letter

        for i, w in enumerate(widths or [], start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "B2"
        return ws

    rows = []
    for s in samples:
        r = deconv.fit(s)
        for p in r.populations:
            rows.append([s.label, "популяция", p.kind, round(p.weight_pct, 2), round(p.mode_um, 3),
                         p.lo_um and round(p.lo_um, 2), p.hi_um and round(p.hi_um, 2), None, round(r.r2, 5), r.k])
        for c in r.components:
            rows.append([s.label, "компонента", c.kind, round(c.weight_pct, 2), None, None, None,
                         round(c.median_um, 3), round(c.sigma_g, 3), None])
    sheet("Популяции", ["Образец", "Строка", "Тип", "Доля, %", "Мода, мкм", "От, мкм", "До, мкм",
                        "Медиана, мкм", "σg / R²", "K (BIC)"], rows, [26, 12, 14, 9, 10, 9, 9, 12, 10, 8])

    wins = [("СЛС", tuple(w)) for w in st.sls_windows] + [("СЭЛС", tuple(w)) for w in st.ebm_windows]
    sheet("Окна печати", ["Образец"] + [f"{t} {w[0]:g}–{w[1]:g}, %" for t, w in wins],
          [[s.label] + [round(windows.frac(s, *w), 2) for _, w in wins] for s in samples], [26] + [14] * len(wins))

    lo, hi = st.sieve_window
    rows = []
    for s in samples:
        sv = windows.sieve(s, lo, hi)
        b = windows.requirements_check(s, st.requirements)
        a = windows.requirements_check(sv.sieved, st.requirements)
        rows.append([s.label, f"{lo:g}–{hi:g}", round(sv.yield_pct, 2), round(sv.fines_pct, 2),
                     round(sv.coarse_pct, 2), round(sv.grams_per_kg(), 0)]
                    + [round(x.value, 2) for x in b] + [round(x.value, 2) for x in a]
                    + [f"{sum(x.ok for x in b)}/{len(b)}", f"{sum(x.ok for x in a)}/{len(a)}"])
    names = [x.name for x in windows.requirements_check(samples[0], st.requirements)] if samples else []
    sheet("Выход годного", ["Образец", "Окно, мкм", "Годное, %", "Мелочь, %", "Крупное, %", "Из 1 кг, г"]
          + [f"исх. {n}" for n in names] + [f"после рассева {n}" for n in names]
          + ["Требования: исх.", "после рассева"], rows, [26, 10] + [11] * 14)

    rho = st.density_g_cm3
    parts = surface.partition(st.windows)
    head = ["Образец", "D[3,2], мкм", "SSA, м²/г"]
    for lo_, hi_ in parts:
        lab = window_label(lo_, hi_)
        head += [f"{lab}: объём, %", f"{lab}: поверхность, %"]
    rows = []
    for s in samples:
        r = [s.label, round(d32(s), 3), round(surface.ssa_m2_g(s, rho), 4)]
        for row in surface.shares(s, st.windows):
            r += [round(row.volume_pct, 2), round(row.surface_pct, 2)]
        rows.append(r)
    ws = sheet("Поверхность", head, rows, [26, 11, 10] + [13] * (len(head) - 3))
    ws.append([])
    ws.append([f"ρ = {rho:g} г/см³. " + surface.ASSUMPTIONS])
