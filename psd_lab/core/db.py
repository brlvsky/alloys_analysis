"""М1. База данных «структура — свойства» (SQLite, без ORM).

Таблицы: batches, measurements, psd_points, psd_metrics, sem, xrd, chem, print_jobs, mech_tests.
Повторный импорт того же измерения не дублирует запись: ключ (file_sha1, meas_no).
"""
from __future__ import annotations

import datetime as dt
import json
import re
import sqlite3
from dataclasses import dataclass
from pathlib import Path

import numpy as np

from .model import Sample

SCHEMA = """
CREATE TABLE IF NOT EXISTS batches(
    id INTEGER PRIMARY KEY, name TEXT UNIQUE NOT NULL, alloy TEXT, additive TEXT, additive_wt_pct REAL,
    state TEXT, route TEXT, composition TEXT, density_measured REAL, notes TEXT);
CREATE TABLE IF NOT EXISTS measurements(
    id INTEGER PRIMARY KEY, batch_id INTEGER NOT NULL REFERENCES batches(id) ON DELETE CASCADE,
    instrument TEXT, meas_no TEXT NOT NULL, file TEXT, file_sha1 TEXT NOT NULL, sheet TEXT, date TEXT,
    model TEXT, obscuration REAL, error REAL, tradeoff REAL, ultrasonics REAL, pump REAL,
    flags TEXT, qc_checklist TEXT, imported_at TEXT, UNIQUE(file_sha1, meas_no));
CREATE TABLE IF NOT EXISTS psd_points(
    measurement_id INTEGER NOT NULL REFERENCES measurements(id) ON DELETE CASCADE,
    size_um REAL NOT NULL, cum_pct REAL NOT NULL);
CREATE INDEX IF NOT EXISTS ix_points ON psd_points(measurement_id);
CREATE TABLE IF NOT EXISTS psd_metrics(
    measurement_id INTEGER PRIMARY KEY REFERENCES measurements(id) ON DELETE CASCADE,
    d10 REAL, d50 REAL, d90 REAL, span REAL, d43 REAL, d32 REAL, ssa_m2_g REAL, fine_pop_pct REAL,
    fractions TEXT, modules TEXT);
CREATE TABLE IF NOT EXISTS sem(
    id INTEGER PRIMARY KEY, batch_id INTEGER NOT NULL REFERENCES batches(id) ON DELETE CASCADE,
    image_path TEXT, magnification REAL, notes TEXT);
CREATE TABLE IF NOT EXISTS xrd(
    id INTEGER PRIMARY KEY, batch_id INTEGER NOT NULL REFERENCES batches(id) ON DELETE CASCADE,
    file_path TEXT, phases TEXT, notes TEXT);
CREATE TABLE IF NOT EXISTS chem(
    id INTEGER PRIMARY KEY, batch_id INTEGER NOT NULL REFERENCES batches(id) ON DELETE CASCADE,
    O_ppm REAL, N_ppm REAL, H_ppm REAL, other TEXT);
CREATE TABLE IF NOT EXISTS print_jobs(
    id INTEGER PRIMARY KEY, batch_id INTEGER NOT NULL REFERENCES batches(id) ON DELETE CASCADE,
    process TEXT, machine TEXT, power_W REAL, speed_mm_s REAL, hatch_um REAL, layer_um REAL, strategy TEXT,
    preheat_C REAL, energy_density_J_mm3 REAL, rel_density_pct REAL);
CREATE TABLE IF NOT EXISTS mech_tests(
    id INTEGER PRIMARY KEY, print_job_id INTEGER NOT NULL REFERENCES print_jobs(id) ON DELETE CASCADE,
    test_type TEXT, temp_C REAL, uts_MPa REAL, ys_MPa REAL, elong_pct REAL, cycles INTEGER, notes TEXT);
"""

TODO_NOTE = "уточнить: смысл обозначения состояния (N/C, П/С) и маршрут получения порошка — вопрос к руководителю"


# ==================================================================== описание полей (для форм)
@dataclass
class Field:
    key: str
    label: str
    kind: str = "text"            # text / memo / float / int / choice / file / image / printjob
    choices: tuple = ()
    required: bool = False


FIELDS: dict[str, list[Field]] = {
    "batches": [
        Field("name", "Название партии", required=True),
        Field("alloy", "Сплав"),
        Field("state", "Состояние (N/C, П/С…)"),
        Field("additive", "Добавка", "choice", ("", "Si", "C", "Y2O3")),
        Field("additive_wt_pct", "Добавка, мас.%", "float"),
        Field("route", "Маршрут получения"),
        Field("density_measured", "Измеренная плотность, г/см³", "float"),
        Field("notes", "Примечания", "memo"),
    ],
    "sem": [Field("image_path", "Файл снимка", "image"), Field("magnification", "Увеличение, ×", "float"),
            Field("notes", "Примечания", "memo")],
    "xrd": [Field("file_path", "Файл дифрактограммы", "file"), Field("phases", "Фазы"),
            Field("notes", "Примечания", "memo")],
    "chem": [Field("O_ppm", "O, ppm", "float"), Field("N_ppm", "N, ppm", "float"), Field("H_ppm", "H, ppm", "float"),
             Field("other", "Другое")],
    "print_jobs": [
        Field("process", "Процесс", "choice", ("СЛС", "СЭЛС")), Field("machine", "Установка"),
        Field("power_W", "Мощность P, Вт", "float"), Field("speed_mm_s", "Скорость v, мм/с", "float"),
        Field("hatch_um", "Шаг штриховки h, мкм", "float"), Field("layer_um", "Толщина слоя t, мкм", "float"),
        Field("strategy", "Стратегия сканирования"), Field("preheat_C", "Подогрев, °C", "float"),
        Field("rel_density_pct", "Относительная плотность, %", "float"),
    ],
    "mech_tests": [
        Field("print_job_id", "Режим печати", "printjob", required=True),
        Field("test_type", "Испытание", "choice", ("растяжение", "сжатие", "усталость", "жаростойкость")),
        Field("temp_C", "Температура, °C", "float"), Field("uts_MPa", "σв, МПа", "float"),
        Field("ys_MPa", "σ0,2, МПа", "float"), Field("elong_pct", "δ, %", "float"),
        Field("cycles", "Число циклов", "int"), Field("notes", "Примечания", "memo"),
    ],
}
TABLE_TITLES = {"sem": "СЭМ", "xrd": "РФА", "chem": "Химанализ", "print_jobs": "Печать", "mech_tests": "Механика"}


# ==================================================================== соединение
def connect(path) -> sqlite3.Connection:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(str(path), timeout=30)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    conn.executescript(SCHEMA)
    conn.commit()
    return conn


def db_path(data_dir: Path) -> Path:
    return Path(data_dir) / "psd.sqlite"


# ==================================================================== партии
ADDITIVES = {"si": "Si", "c": "C", "y2o3": "Y2O3"}


def parse_name(name: str) -> dict:
    """«П/С +0,5Y2O3» → state П/С, additive Y2O3, 0.5 мас.%; «N/C+si» → N/C, Si."""
    out = {"state": None, "additive": None, "additive_wt_pct": None, "alloy": None, "notes": None}
    m = re.match(r"^\s*(N/C|П/С|NC|ПС)\s*(.*)$", name, flags=re.IGNORECASE)
    if not m:
        h = re.match(r"^\s*(\d+(?:[.,]\d+)?)\s*(ч|h|мин|min)\s*$", name, flags=re.IGNORECASE)
        if h:
            out["notes"] = f"время обработки {h.group(1)} {h.group(2)}; материал неизвестен (старый шаблон «Расчет.xlsx»)"
        return out
    out["state"] = m.group(1).upper().replace("NC", "N/C").replace("ПС", "П/С")
    out["alloy"] = "TANMB (Ti–Al–Nb–Mo–B, γ-TiAl; состав уточнить)"
    out["notes"] = TODO_NOTE
    rest = m.group(2).strip().lstrip("+").strip()
    a = re.match(r"^(\d+(?:[.,]\d+)?)?\s*([A-Za-z][A-Za-z0-9]*)$", rest)
    if a:
        out["additive"] = ADDITIVES.get(a.group(2).lower(), a.group(2))
        if a.group(1):
            out["additive_wt_pct"] = float(a.group(1).replace(",", "."))
    return out


def get_or_create_batch(conn, name: str) -> int:
    row = conn.execute("SELECT id FROM batches WHERE name = ?", (name,)).fetchone()
    if row:
        return row["id"]
    p = parse_name(name)
    cur = conn.execute("INSERT INTO batches(name, alloy, additive, additive_wt_pct, state, notes) VALUES (?,?,?,?,?,?)",
                       (name, p["alloy"], p["additive"], p["additive_wt_pct"], p["state"], p["notes"]))
    return cur.lastrowid


# ==================================================================== импорт измерений
def meas_key(s: Sample) -> str:
    """Номер измерения: M#### для прибора; лист/столбец для таблиц."""
    if s.meas_id:
        return s.meas_id
    return f"{s.sheet}/{s.meta.get('column', s.name)}"


def exists(conn, sha1: str, s: Sample) -> bool:
    return conn.execute("SELECT 1 FROM measurements WHERE file_sha1=? AND meas_no=?",
                        (sha1, meas_key(s))).fetchone() is not None


def module_results(s: Sample, settings=None) -> dict:
    """Результаты модулей М2–М5 для хранения в psd_metrics.modules (JSON)."""
    from . import deconv, surface, windows

    st = settings
    sls_w = st.sls_windows if st else [[15, 45], [15, 53], [20, 63]]
    ebm_w = st.ebm_windows if st else [[45, 105], [45, 150]]
    lo, hi = st.sieve_window if st else (15, 53)
    r = deconv.fit(s)
    sv = windows.sieve(s, lo, hi)
    v15, a15 = surface.fine_shares(s)
    return {
        "deconv": {"k": r.k, "r2": round(r.r2, 6),
                   "populations": [{"kind": p.kind, "weight_pct": round(p.weight_pct, 3),
                                    "mode_um": round(p.mode_um, 4)} for p in r.populations],
                   "components": [{"kind": c.kind, "weight_pct": round(c.weight_pct, 3),
                                   "median_um": round(c.median_um, 4), "sigma_g": round(c.sigma_g, 4)}
                                  for c in r.components]},
        "windows": {f"СЛС {a:g}-{b:g}": round(windows.frac(s, a, b), 3) for a, b in sls_w}
        | {f"СЭЛС {a:g}-{b:g}": round(windows.frac(s, a, b), 3) for a, b in ebm_w},
        "sieve": {"window": [lo, hi], "yield_pct": round(sv.yield_pct, 3)},
        "surface": {"lt15_volume_pct": round(v15, 3), "lt15_surface_pct": round(a15, 3)},
    }


def import_sample(conn, s: Sample, sha1: str, settings=None, modules=True) -> int | None:
    """Импорт одного измерения. Возвращает id новой записи или None, если оно уже в базе."""
    from .metrics import compute
    from .surface import ssa_m2_g

    if exists(conn, sha1, s):
        return None
    bid = get_or_create_batch(conn, s.name)
    m = s.meta
    date = m.get("date")
    cur = conn.execute(
        "INSERT INTO measurements(batch_id, instrument, meas_no, file, file_sha1, sheet, date, model, obscuration, "
        "error, tradeoff, ultrasonics, pump, flags, imported_at) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
        (bid, "Fritsch ANALYSETTE 22" if s.source == "fritsch" else "таблица", meas_key(s), s.file, sha1, s.sheet,
         date.isoformat(sep=" ", timespec="minutes") if hasattr(date, "isoformat") else date, m.get("model"),
         m.get("obscuration"), m.get("error"), m.get("tradeoff"), m.get("ultrasonics"), m.get("pump"),
         json.dumps([list(f) for f in s.flags], ensure_ascii=False),
         dt.datetime.now().isoformat(sep=" ", timespec="seconds")))
    mid = cur.lastrowid
    conn.executemany("INSERT INTO psd_points(measurement_id, size_um, cum_pct) VALUES (?,?,?)",
                     [(mid, float(x), float(c)) for x, c in zip(s.size_um, s.cum_pct)])
    windows_ = settings.windows_tuples if settings else None
    mm = compute(s, windows_)
    rho = settings.density_g_cm3 if settings else 4.0
    mods = module_results(s, settings) if modules else None
    fine = None
    if mods:
        pops = [p for p in mods["deconv"]["populations"] if p["kind"] != "субмикронная"]
        fine = pops[0]["weight_pct"] if len(pops) >= 2 else None
    conn.execute(
        "INSERT INTO psd_metrics(measurement_id, d10, d50, d90, span, d43, d32, ssa_m2_g, fine_pop_pct, fractions, "
        "modules) VALUES (?,?,?,?,?,?,?,?,?,?,?)",
        (mid, mm["d10"], mm["d50"], mm["d90"], mm["span"], mm["d43"], mm["d32"], ssa_m2_g(s, rho), fine,
         json.dumps({k: round(v, 4) for k, v in mm["fractions"].items()}, ensure_ascii=False),
         json.dumps(mods, ensure_ascii=False) if mods else None))
    return mid


def import_samples(conn, samples: list[Sample], sha1: str, settings=None, modules=True) -> int:
    """Импорт всех измерений файла (исходных, не усреднённых). Возвращает число новых записей."""
    n = 0
    with conn:
        for s in samples:
            if import_sample(conn, s, sha1, settings, modules) is not None:
                n += 1
    return n


def load_curve(conn, measurement_id: int) -> tuple[np.ndarray, np.ndarray]:
    rows = conn.execute("SELECT size_um, cum_pct FROM psd_points WHERE measurement_id=? ORDER BY size_um",
                        (measurement_id,)).fetchall()
    return np.array([r[0] for r in rows]), np.array([r[1] for r in rows])


# ==================================================================== чтение и правка таблиц
def batches(conn) -> list[sqlite3.Row]:
    return conn.execute(
        "SELECT b.*, (SELECT COUNT(*) FROM measurements m WHERE m.batch_id=b.id) AS n_meas "
        "FROM batches b ORDER BY b.name").fetchall()


def batch_measurements(conn, batch_id: int) -> list[sqlite3.Row]:
    return conn.execute(
        "SELECT m.*, p.d10, p.d50, p.d90, p.span, p.d43, p.d32, p.ssa_m2_g, p.fine_pop_pct, p.fractions "
        "FROM measurements m LEFT JOIN psd_metrics p ON p.measurement_id = m.id WHERE m.batch_id=? "
        "ORDER BY m.date, m.meas_no", (batch_id,)).fetchall()


def rows(conn, table: str, batch_id: int) -> list[sqlite3.Row]:
    if table == "mech_tests":
        return conn.execute(
            "SELECT t.*, j.process || ' ' || COALESCE(j.machine, '') || ' #' || j.id AS job "
            "FROM mech_tests t JOIN print_jobs j ON j.id = t.print_job_id WHERE j.batch_id=? ORDER BY t.id",
            (batch_id,)).fetchall()
    return conn.execute(f"SELECT * FROM {_tbl(table)} WHERE batch_id=? ORDER BY id", (batch_id,)).fetchall()


def get(conn, table: str, rid: int):
    return conn.execute(f"SELECT * FROM {_tbl(table)} WHERE id=?", (rid,)).fetchone()


def _tbl(table: str) -> str:
    if table not in FIELDS:
        raise ValueError(f"неизвестная таблица {table}")
    return table


def energy_density(power_W, speed_mm_s, hatch_um, layer_um) -> float | None:
    """Объёмная плотность энергии E = P / (v · h · t), Дж/мм³ (h и t переводятся из мкм в мм)."""
    try:
        if not all(v and v > 0 for v in (power_W, speed_mm_s, hatch_um, layer_um)):
            return None
        return power_W / (speed_mm_s * (hatch_um / 1000) * (layer_um / 1000))
    except TypeError:
        return None


def coerce(table: str, values: dict) -> dict:
    """Проверка и приведение типов по описанию полей. Ошибки — ValueError с понятным текстом."""
    out = {}
    for f in FIELDS[table]:
        v = values.get(f.key)
        if isinstance(v, str):
            v = v.strip()
        if v in ("", None):
            if f.required:
                raise ValueError(f"Заполните поле «{f.label}».")
            out[f.key] = None
            continue
        if f.kind in ("float", "int", "printjob"):
            try:
                v = float(str(v).replace(",", "."))
            except ValueError:
                raise ValueError(f"«{f.label}»: нужно число.") from None
            if f.kind in ("int", "printjob"):
                v = int(round(v))
        out[f.key] = v
    if table == "print_jobs":
        out["energy_density_J_mm3"] = energy_density(out.get("power_W"), out.get("speed_mm_s"),
                                                     out.get("hatch_um"), out.get("layer_um"))
    return out


def insert(conn, table: str, values: dict, batch_id: int | None = None) -> int:
    v = coerce(table, values)
    if table not in ("batches", "mech_tests"):
        v["batch_id"] = batch_id
    cols = ", ".join(v)
    with conn:
        cur = conn.execute(f"INSERT INTO {_tbl(table)}({cols}) VALUES ({', '.join('?' * len(v))})", list(v.values()))
    return cur.lastrowid


def update(conn, table: str, rid: int, values: dict) -> None:
    v = coerce(table, values)
    with conn:
        conn.execute(f"UPDATE {_tbl(table)} SET {', '.join(f'{k}=?' for k in v)} WHERE id=?", [*v.values(), rid])


def save_composition(conn, batch_id: int, composition: dict, measured: float | None = None) -> None:
    """Состав из калькулятора плотности (М7) — в карточку партии; измеренная плотность — если задана."""
    with conn:
        conn.execute("UPDATE batches SET composition=? WHERE id=?",
                     (json.dumps(composition, ensure_ascii=False), batch_id))
        if measured:
            conn.execute("UPDATE batches SET density_measured=? WHERE id=?", (float(measured), batch_id))


def delete(conn, table: str, rid: int) -> None:
    with conn:
        conn.execute(f"DELETE FROM {_tbl(table)} WHERE id=?", (rid,))


def delete_measurement(conn, mid: int) -> None:
    with conn:
        conn.execute("DELETE FROM measurements WHERE id=?", (mid,))


# ==================================================================== структура — свойства
PROPERTY_FIELDS: list[tuple[str, str]] = [
    ("psd.d10", "d10, мкм"), ("psd.d50", "d50, мкм"), ("psd.d90", "d90, мкм"), ("psd.span", "span"),
    ("psd.d43", "D[4,3], мкм"), ("psd.d32", "D[3,2], мкм"), ("psd.ssa_m2_g", "Удельная поверхность, м²/г"),
    ("psd.fine_pop_pct", "Мелкая популяция, %"),
    ("batch.additive_wt_pct", "Добавка, мас.%"), ("batch.density_measured", "Измеренная плотность, г/см³"),
    ("chem.O_ppm", "O, ppm"), ("chem.N_ppm", "N, ppm"), ("chem.H_ppm", "H, ppm"),
    ("print.power_W", "Мощность, Вт"), ("print.speed_mm_s", "Скорость, мм/с"),
    ("print.energy_density_J_mm3", "Плотность энергии, Дж/мм³"), ("print.preheat_C", "Подогрев, °C"),
    ("print.rel_density_pct", "Относительная плотность, %"),
    ("mech.uts_MPa", "σв, МПа"), ("mech.ys_MPa", "σ0,2, МПа"), ("mech.elong_pct", "δ, %"),
]


def property_fields(conn) -> list[tuple[str, str]]:
    """Все числовые поля, доступные для графика «структура — свойства» (с долями по окнам из базы)."""
    out = list(PROPERTY_FIELDS)
    labels = set()
    for (fr,) in conn.execute("SELECT fractions FROM psd_metrics WHERE fractions IS NOT NULL"):
        labels |= set(json.loads(fr))
    for lab in sorted(labels, key=_label_order):
        out.insert(8, (f"frac.{lab}", f"Доля {lab} мкм, %"))
    return out


def _label_order(lab: str):
    nums = re.findall(r"\d+", lab)
    return (0 if lab.startswith("<") else 2 if lab.startswith(">") else 1, float(nums[0]) if nums else 0)


def property_values(conn, key: str) -> dict[str, float]:
    """Значение поля по партиям (среднее, если записей несколько)."""
    src, col = key.split(".", 1)
    q = {
        "psd": f"SELECT b.name, p.{col} FROM psd_metrics p JOIN measurements m ON m.id=p.measurement_id "
               "JOIN batches b ON b.id=m.batch_id",
        "batch": f"SELECT name, {col} FROM batches",
        "chem": f"SELECT b.name, c.{col} FROM chem c JOIN batches b ON b.id=c.batch_id",
        "print": f"SELECT b.name, j.{col} FROM print_jobs j JOIN batches b ON b.id=j.batch_id",
        "mech": f"SELECT b.name, t.{col} FROM mech_tests t JOIN print_jobs j ON j.id=t.print_job_id "
                "JOIN batches b ON b.id=j.batch_id",
    }
    acc: dict[str, list[float]] = {}
    if src == "frac":
        for name, fr in conn.execute("SELECT b.name, p.fractions FROM psd_metrics p JOIN measurements m "
                                     "ON m.id=p.measurement_id JOIN batches b ON b.id=m.batch_id"):
            v = json.loads(fr or "{}").get(col)
            if v is not None:
                acc.setdefault(name, []).append(v)
    else:
        if src not in q or not re.fullmatch(r"[A-Za-z0-9_]+", col):
            raise ValueError(f"неизвестное поле {key}")
        for name, v in conn.execute(q[src]):
            if v is not None:
                acc.setdefault(name, []).append(float(v))
    return {k: float(np.mean(v)) for k, v in acc.items()}


@dataclass
class Regression:
    n: int
    slope: float | None
    intercept: float | None
    r2: float | None          # только при n ≥ 3


def xy(conn, kx: str, ky: str) -> list[tuple[str, float, float]]:
    vx, vy = property_values(conn, kx), property_values(conn, ky)
    return sorted((b, vx[b], vy[b]) for b in vx.keys() & vy.keys())


def regression(x, y) -> Regression:
    x, y = np.asarray(x, float), np.asarray(y, float)
    n = len(x)
    if n < 2 or np.ptp(x) == 0:
        return Regression(n, None, None, None)
    k, b = np.polyfit(x, y, 1)
    r2 = None
    if n >= 3:
        ss_res = float(np.sum((y - (k * x + b)) ** 2))
        ss_tot = float(np.sum((y - y.mean()) ** 2))
        r2 = 1 - ss_res / ss_tot if ss_tot > 0 else None
    return Regression(n, float(k), float(b), r2)


# ==================================================================== экспорт
EXPORT_TABLES = [("batches", "Партии"), ("measurements", "Измерения"), ("psd_metrics", "Метрики"),
                 ("psd_points", "Точки кривых"), ("sem", "СЭМ"), ("xrd", "РФА"), ("chem", "Химанализ"),
                 ("print_jobs", "Печать"), ("mech_tests", "Механика")]


def export_xlsx(conn, path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    wb.remove(wb.active)
    for table, title in EXPORT_TABLES:
        cur = conn.execute(f"SELECT * FROM {table}")
        cols = [d[0] for d in cur.description]
        ws = wb.create_sheet(title)
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill = Font(bold=True), PatternFill("solid", fgColor="D9D9D9")
        for r in cur:
            ws.append(list(r))
        ws.freeze_panes = "A2"
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
