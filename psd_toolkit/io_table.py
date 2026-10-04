"""Чтение произвольных таблиц (.xlsx, .xls, .csv, .txt) и общий читатель книг Excel.

Логика: на каждом листе ищем столбец размеров (самая длинная строго возрастающая
неотрицательная числовая последовательность, не короче 5 значений), затем рядом —
столбцы распределения (накопленные или по интервалам).
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import re
from pathlib import Path

import numpy as np

from .model import Sample

MIN_POINTS = 5
GENERIC_HEADERS = {
    "freqcum", "q", "q,%", "q, %", "%", "σq", "σq,%", "σq, %", "q3", "cum", "cumulative",
    "sumq", "накопл", "накопленная", "доля", "value", "values", "y", "tt",
}
GENERIC_SHEETS = {"лист1", "sheet1", "лист", "sheet", "data", "данные"}


# ---------------------------------------------------------------- чтение файлов
def read_grids(path: Path) -> list[tuple[str, list[list]]]:
    """Возвращает [(имя листа, строки)], ячейки — float / str / datetime / None."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".xls":
        return _read_xls(path)
    if ext in (".xlsx", ".xlsm"):
        return _read_xlsx(path)
    if ext in (".csv", ".txt", ".tsv", ".dat"):
        return [(path.stem, _read_text(path))]
    raise ValueError(f"Неизвестный формат файла: {path.name}")


def _read_xls(path: Path):
    import xlrd

    book = xlrd.open_workbook(str(path))
    out = []
    for sh in book.sheets():
        rows = []
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    row.append(None)
                elif cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                elif cell.ctype in (xlrd.XL_CELL_NUMBER, xlrd.XL_CELL_BOOLEAN):
                    row.append(float(cell.value))
                elif cell.ctype == xlrd.XL_CELL_TEXT:
                    row.append(cell.value if cell.value.strip() else None)
                else:
                    row.append(None)
            rows.append(row)
        out.append((sh.name, rows))
    return out


def _read_xlsx(path: Path):
    import openpyxl

    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = []
            for v in row:
                if isinstance(v, bool):
                    v = float(v)
                elif isinstance(v, (int, float)):
                    v = float(v)
                elif isinstance(v, str) and not v.strip():
                    v = None
                cells.append(v)
            rows.append(cells)
        out.append((ws.title, rows))
    wb.close()
    return out


def _read_text(path: Path) -> list[list]:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1251", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    lines = [ln for ln in text.splitlines() if ln.strip()]
    sample = "\n".join(lines[:50])
    if "\t" in sample:
        delim = "\t"
    elif ";" in sample:
        delim = ";"
    elif "," in sample and not _whitespace_table(lines[:50]):
        delim = ","  # запятая-разделитель, десятичная точка
    else:
        delim = None  # пробелы (числа могут быть с десятичной запятой)
    decimal_comma = delim != ","
    rows = []
    if delim is None:
        parsed = [ln.split() for ln in lines]
    else:
        parsed = list(csv.reader(io.StringIO("\n".join(lines)), delimiter=delim))
    for parts in parsed:
        rows.append([_parse_cell(p, decimal_comma) for p in parts])
    return rows


def _whitespace_table(lines) -> bool:
    """Строки вида '1,5  2,3' — столбцы через пробелы, десятичная запятая."""
    data = [ln.split() for ln in lines if re.match(r"\s*[\d.,]", ln)]
    return bool(data) and all(len(t) >= 2 and all(re.fullmatch(r"-?\d+(,\d+)?", x) for x in t) for t in data)


def _parse_cell(s: str, decimal_comma: bool):
    s = s.strip().strip('"')
    if not s:
        return None
    t = s.replace(" ", "").replace(" ", "")
    if decimal_comma:
        t = t.replace(",", ".")
    try:
        return float(t)
    except ValueError:
        return s


# ---------------------------------------------------------------- разбор таблиц
def _num(v):
    return isinstance(v, float) and np.isfinite(v)


def _column(rows, c):
    return [row[c] if c < len(row) else None for row in rows]


def _longest_increasing_run(col) -> tuple[int, int]:
    """(начало, длина) самой длинной строго возрастающей неотрицательной числовой серии."""
    best = (0, 0)
    start, length, prev = 0, 0, None
    for i, v in enumerate(col):
        if _num(v) and v >= 0 and (prev is None or v > prev):
            if length == 0:
                start = i
            length += 1
            prev = v
        else:
            if length > best[1]:
                best = (start, length)
            if _num(v) and v >= 0:
                start, length, prev = i, 1, v
            else:
                length, prev = 0, None
    if length > best[1]:
        best = (start, length)
    return best


def _header_above(rows, c, r0) -> str | None:
    for r in range(r0 - 1, max(-1, r0 - 6), -1):
        v = rows[r][c] if c < len(rows[r]) else None
        if isinstance(v, str) and v.strip():
            return v.strip()
        if v is not None and not isinstance(v, str):
            return None
    return None


def _classify(vals: np.ndarray):
    """'cum' | 'int' | None, и значения в процентах."""
    if np.nanmax(vals) <= 1.0 + 1e-9:
        vals = vals * 100.0
    if np.all(np.diff(vals) >= -1e-9) and 90.0 <= vals[-1] <= 101.0:
        return "cum", vals
    if np.all(vals >= -1e-9) and 90.0 <= vals.sum() <= 105.0:
        return "int", vals
    return None, vals


def parse_grid(rows: list[list], sheet: str, file: Path) -> list[Sample]:
    ncols = max((len(r) for r in rows), default=0)
    best_c, best = None, (0, 0)
    for c in range(ncols):
        run = _longest_increasing_run(_column(rows, c))
        if run[1] > best[1]:
            best_c, best = c, run
    if best_c is None or best[1] < MIN_POINTS:
        return []
    r0, n = best
    sizes = np.array(_column(rows, best_c)[r0:r0 + n], dtype=float)
    found = {"cum": [], "int": []}
    for c in range(ncols):
        if c == best_c:
            continue
        col = _column(rows, c)[r0:r0 + n]
        if not all(_num(v) for v in col):  # короткие столбцы (таблица по 5 мкм и т.п.) — мимо
            continue
        kind, vals = _classify(np.array(col, dtype=float))
        if kind:
            found[kind].append((c, vals))
    kind = "cum" if found["cum"] else "int"
    out = []
    for c, vals in found[kind]:
        cum = vals if kind == "cum" else np.cumsum(vals)
        header = _header_above(rows, c, r0)
        out.append(Sample(
            name=_sample_name(header, sheet, file),
            size_um=sizes, cum_pct=cum, file=str(file), sheet=sheet, source="table",
            meta={"column": header or f"#{c + 1}", "distribution": kind},
        ))
    return out


def _sample_name(header, sheet, file: Path) -> str:
    if header and header.strip().lower().replace(" ", "") not in {h.replace(" ", "") for h in GENERIC_HEADERS}:
        return header
    if sheet and sheet.strip().lower() not in GENERIC_SHEETS and sheet != Path(file).stem:
        return sheet
    return Path(file).stem


def load_table(path: Path) -> list[Sample]:
    path = Path(path)
    samples = []
    for sheet, rows in read_grids(path):
        samples.extend(parse_grid(rows, sheet, path))
    return samples


def excel_serial_to_datetime(v) -> dt.datetime | None:
    if isinstance(v, dt.datetime):
        return v
    if _num(v):
        return dt.datetime(1899, 12, 30) + dt.timedelta(days=v)
    return None
